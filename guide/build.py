"""Build the 'How to Generate AccuKnox-Branded Docs and PPTs with Claude Code' user guide.

Updated for the new standalone doc-ppt-template repo (templates now live at repo root,
not inside a Help Docs subfolder). Images are reused as-is from the original guide.
"""
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
IMAGES = os.path.join(HERE, "images")
LOGO = os.path.join(IMAGES, "logo.png")
OUT = os.path.join(HERE, "How_to_Generate_AccuKnox_Branded_Docs_and_PPTs_with_Claude_Code.docx")

NAVY = RGBColor(0x11, 0x20, 0x6D)
ACCENT = RGBColor(0x4D, 0x4D, 0xD9)
DARK = RGBColor(0x0D, 0x1B, 0x4B)
MID = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xC8, 0x00, 0x19)

HEAD_FONT = "Space Grotesk"
BODY_FONT = "Inter"


def set_margins(doc, top=0.6, bottom=0.6, left=0.7, right=0.7):
    for section in doc.sections:
        section.top_margin = Inches(top)
        section.bottom_margin = Inches(bottom)
        section.left_margin = Inches(left)
        section.right_margin = Inches(right)


def add_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def para(doc, text="", size=11, color=DARK, bold=False, italic=False, font=BODY_FONT,
          space_before=0, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    return p


def heading(doc, text, size=16, color=NAVY, space_before=10, space_after=6):
    return para(doc, text, size=size, color=color, bold=True, font=HEAD_FONT,
                space_before=space_before, space_after=space_after)


def numbered(doc, num, text, color=DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.05)
    run = p.add_run(f"{num}.  ")
    run.font.bold = True
    run.font.color.rgb = ACCENT
    run.font.size = Pt(11)
    run.font.name = BODY_FONT
    run2 = p.add_run(text)
    run2.font.size = Pt(11)
    run2.font.color.rgb = color
    run2.font.name = BODY_FONT
    return p


def bullet(doc, text, color=DARK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = color
    run.font.name = BODY_FONT
    return p


def add_image_framed(doc, path, width_in, caption=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        cr = cp.add_run(caption)
        cr.font.size = Pt(9)
        cr.font.italic = True
        cr.font.color.rgb = MID
        cr.font.name = BODY_FONT


doc = Document()
set_margins(doc)

# Base style
normal = doc.styles["Normal"]
normal.font.name = BODY_FONT
normal.font.size = Pt(11)
normal.font.color.rgb = DARK

# --- Header: logo + title ---
htab = doc.add_table(rows=1, cols=2)
htab.alignment = WD_TABLE_ALIGNMENT.CENTER
htab.autofit = True
c0, c1 = htab.rows[0].cells
c0.width = Inches(1.3)
c1.width = Inches(5.9)
p0 = c0.paragraphs[0]
r0 = p0.add_run()
if os.path.exists(LOGO):
    r0.add_picture(LOGO, width=Inches(1.1))
p1 = c1.paragraphs[0]
p1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
r1 = p1.add_run("How to Generate AccuKnox-Branded\nDocs and PPTs with Claude Code")
r1.font.size = Pt(18)
r1.font.bold = True
r1.font.color.rgb = NAVY
r1.font.name = HEAD_FONT

para(doc, "", size=4)

# --- Intro ---
para(doc,
     "Claude Code can take any Word doc or Google Slides deck and reformat it into AccuKnox "
     "branding, correct logo, colors, and layout, in one pass. Works for both docs and PPTs, "
     "and you can batch as many files as you want in one go.",
     size=11, color=DARK, space_after=10)

# --- Prerequisites ---
heading(doc, "Prerequisites", size=14, space_before=4)
bullet(doc, "An active Claude Code subscription. We use one shared account "
            "(website@accuknox.com, 30+ users). Claude desktop must be installed, and you "
            "work inside the Claude Code window.")
bullet(doc, "The docs or PPTs you want branded, saved locally on your PC.")
bullet(doc, "The doc-ppt-template repo cloned locally: "
            "https://github.com/accuknox/doc-ppt-template")

para(doc, "", size=4)

# --- Steps ---
heading(doc, "Steps", size=14, space_before=4)
numbered(doc, 1, "Clone the doc-ppt-template repo and load it in Claude Code.")
numbered(doc, 2, "Attach the docx or PPT you want to rebrand into AccuKnox's branding.")
numbered(doc, 3, "Reference the template files using the @ selector. Both live at the root "
                 "of the doc-ppt-template repo:")
bullet(doc, "Word doc → WORD_TEMPLATE_ACCUKNOX.docx")
bullet(doc, "Presentation → PPT Template.pptx (or AccuKnox_Proposal_Template_BLANK.pptx to start a new deck)")
numbered(doc, 4, "Set permissions to “Bypass permissions” so Claude works autonomously.")
numbered(doc, 5, "Use Sonnet 5 only. Do not use Opus for this task.")

para(doc, "", size=4)

# --- Example prompt ---
heading(doc, "Example Prompt", size=13, space_before=2, space_after=4)
tbl = doc.add_table(rows=1, cols=1)
cell = tbl.rows[0].cells[0]
add_shading(cell, "F5F6FF")
cp = cell.paragraphs[0]
cp.paragraph_format.space_before = Pt(6)
cp.paragraph_format.space_after = Pt(6)
cr = cp.add_run(
    "@doc-ppt-template Brand the attached DOCX file with AccuKnox branding, do it "
    "well and ensure it looks as per the expected outcome, be strict and use the logo "
    "correctly as well."
)
cr.font.size = Pt(10.5)
cr.font.italic = True
cr.font.color.rgb = NAVY
cr.font.name = BODY_FONT

para(doc, "", size=6)
add_image_framed(doc, os.path.join(IMAGES, "workflow-steps.png"), 6.4,
                  caption="The five steps above, shown inside Claude Code. "
                          "(Screenshot predates the doc-ppt-template repo split; "
                          "the @ selector and Bypass Permissions workflow are unchanged.)")

# --- Before / After ---
heading(doc, "Before / After", size=14, space_before=2)
para(doc, "The same comparison table, before and after AccuKnox branding is applied.",
     size=11, color=MID, space_after=6)
add_image_framed(doc, os.path.join(IMAGES, "before-after.png"), 6.4)

doc.save(OUT)
print("Saved:", OUT)
