"""Insert a YouTube 'Online Video' into a python-pptx slide, mirroring what
PowerPoint 365 (16.0.20326) writes for Insert > Video > Online Video."""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches
from pptx.oxml.ns import qn

RT_VIDEO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video"

TIMING = """<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq"><p:stCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cond></p:stCondLst><p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="2" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:cmd type="call" cmd="togglePause"><p:cBhvr><p:cTn id="6" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:nextCondLst><p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cond></p:nextCondLst></p:seq><p:video><p:cMediaNode vol="80000"><p:cTn id="7" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"""


def add_youtube(slide, video_id, poster_path, left, top, width, height,
                title="", start=None):
    """One video per slide (the timing block assumes it)."""
    url = f"https://www.youtube.com/embed/{video_id}"
    url += f"?start={int(start)}&feature=oembed" if start else "?feature=oembed"
    pic = slide.shapes.add_picture(poster_path, left, top, width, height)
    pic.name = "Online Media"
    rid = slide.part.relate_to(url, RT_VIDEO, is_external=True)
    cNvPr = pic._element.nvPicPr.cNvPr
    if title:
        cNvPr.set("title", title)
        cNvPr.set("descr", f"YouTube video: {title}")  # python-pptx puts the file name here
    h = etree.SubElement(cNvPr, qn("a:hlinkClick"))
    h.set(qn("r:id"), "")
    h.set("action", "ppaction://media")
    # a:hlinkClick must come before a:extLst inside cNvPr
    ext = cNvPr.find(qn("a:extLst"))
    if ext is not None:
        cNvPr.remove(h); ext.addprevious(h)
    cNvPicPr = pic._element.nvPicPr.find(qn("p:cNvPicPr"))
    locks = cNvPicPr.find(qn("a:picLocks"))
    if locks is None:
        locks = etree.SubElement(cNvPicPr, qn("a:picLocks"))
    locks.set("noRot", "1"); locks.set("noChangeAspect", "1")
    nvPr = pic._element.nvPicPr.find(qn("p:nvPr"))
    vf = etree.SubElement(nvPr, qn("a:videoFile"))
    vf.set(qn("r:link"), rid)
    # timing goes after p:clrMapOvr, before p:extLst on p:sld
    sld = slide._element
    old = sld.find(qn("p:timing"))
    if old is not None:
        sld.remove(old)
    timing = etree.fromstring(TIMING.replace("{spid}", str(cNvPr.get("id"))))
    anchor = sld.find(qn("p:extLst"))
    if anchor is not None:
        anchor.addprevious(timing)
    else:
        sld.append(timing)
    return pic


def add_poster_link(slide, video_id, poster_path, left, top, width, height):
    """Fallback: a static poster that opens the video in the browser."""
    pic = slide.shapes.add_picture(poster_path, left, top, width, height)
    pic.click_action.hyperlink.address = f"https://www.youtube.com/watch?v={video_id}"
    pic._element.nvPicPr.cNvPr.set("descr", "Poster image. Click to open the video on YouTube.")
    return pic


if __name__ == "__main__":
    import sys
    thumbs = sys.argv[1]
    out = sys.argv[2]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    add_youtube(s1, "qTDQjmm8698", f"{thumbs}/qTDQjmm8698.jpg",
                Inches(1.5), Inches(0.75), Inches(10.333), Inches(5.8125),
                title="AI Asset Onboarding and Overview")
    s2 = prs.slides.add_slide(blank)
    add_youtube(s2, "l_RCQosnNJk", f"{thumbs}/l_RCQosnNJk.jpg",
                Inches(1.5), Inches(0.75), Inches(10.333), Inches(5.8125),
                title="AccuKnox Prompt Firewall Setup and Demo", start=45)
    s3 = prs.slides.add_slide(blank)
    add_poster_link(s3, "mAzLWcr59g0", f"{thumbs}/mAzLWcr59g0.jpg",
                    Inches(1.5), Inches(0.75), Inches(10.333), Inches(5.8125))
    prs.save(out)
    print("saved", out)
