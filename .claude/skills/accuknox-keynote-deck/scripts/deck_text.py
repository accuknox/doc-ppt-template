# -*- coding: utf-8 -*-
"""deck_text.py: pull every word out of a built deck, for the writing gate.

    py -3.11 deck_text.py output/My_Deck.pptx > work/deck-text.md
    python "D:\\Atharva\\NOTES\\SCRIPTS\\slop\\score.py" work/deck-text.md

Slide text and speaker notes both come out, because the notes carry the source
trail and deserve the same scrutiny as the slide. Two findings are expected in
the scorer output: the `## Slide n` headings read as topic labels, and chip rows
read as verbless fragments, because they are lists rather than sentences.
"""
import sys
from pathlib import Path

from pptx import Presentation


def dump(path):
    prs = Presentation(str(path))
    out = [f"# {Path(path).name}"]
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"\n## Slide {i}\n")
        for sh in slide.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        out.append(t + "\n")
            if sh.has_table:
                for row in sh.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        out.append(" | ".join(cells) + "\n")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append("Notes. " + notes + "\n")
    return "\n".join(out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    for arg in sys.argv[1:]:
        print(dump(arg))
