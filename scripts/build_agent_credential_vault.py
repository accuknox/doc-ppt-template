# -*- coding: utf-8 -*-
"""
One slide: how a credential vault lets AI agents call APIs without holding the keys.
Rebuilt in the AccuKnox visual style from a third-party infographic the team shared.

The flow diagram is an SVG from _aksvg. The six controls and the takeaway are native
shapes, so they stay editable. The full talk track is in the speaker notes.

Build:   py -3.11 scripts/build_agent_credential_vault.py
Render:  powershell -File scripts/render.ps1 -Pptx <abs>/output/AccuKnox_Agent_Credential_Vault.pptx -Out <abs>/output/render/agent_vault
"""
import os, shutil, sys
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _akdeck import *          # noqa: F401,F403
import _aksvg as S

google_safe()

HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(HERE, "..", "PPT Template.pptx")
OUT  = os.path.join(HERE, "..", "output", "AccuKnox_Agent_Credential_Vault.pptx")
DIAG = os.path.join(HERE, "..", "work", "agent_vault", "diagrams")


# =================================================================== diagram
def flow_card(x, y, w, h, icon_name, head, lines, line_icon, line_color, accent=S.PRIMARY,
              fill=S.WHITE, step=None):
    out = [S.rect(x, y, w, h, fill=fill, stroke=accent, sw=2.5),
           f'<circle cx="{x + 36}" cy="{y + 38}" r="23" fill="{S.LAV if fill == S.WHITE else S.WHITE}"/>',
           S.ui(icon_name, x + 20, y + 22, 32, accent),
           S.text(x + 68, y + 45, head, 19, S.NAVY, 700)]
    for i, ln in enumerate(lines):
        ly = y + 100 + i * 44
        out.append(S.ui(line_icon, x + 18, ly - 17, 22, line_color))
        out.append(S.text(x + 48, ly, ln, 17, S.INK, 500))
    if step:
        out.append(S.badge(x + w - 22, y + h - 22, step, r=15, size=17))
    return "".join(out)


def arrow_label(cx, y, lines, n, above=True):
    lh = 21
    ty = y - 14 - lh * (len(lines) - 1) if above else y + 36
    return S.badge(cx, y, n, r=14, size=17) + S.text(cx, ty - (14 if above else 0), lines, 17, S.MUTE, 600,
                                                     anchor="middle", lh=1.2)


def build_flow():
    b = []
    top, h = 10, 220
    # 1. agent, 2. gateway, vault, targets
    b.append(flow_card(0, top, 230, h, "robot", "AI Agent",
                       ["No keys in prompt", "No hardcoded keys", "No standing tokens"],
                       "ban", S.RED))
    b.append(flow_card(330, top, 230, h, "shield-check", "Agent Gateway",
                       ["Verifies identity", "Checks permissions", "Stops risky calls"],
                       "check", S.PRIMARY, step=2))
    b.append(flow_card(660, top, 230, h, "lock-square", "Credential Vault",
                       ["API keys", "OAuth tokens", "Short-lived access"],
                       "key", S.PRIMARY, fill=S.LAV))

    # targets panel
    tx, tw = 980, 300
    b.append(S.rect(tx, top, tw, h, fill=S.GREY_BG, stroke=S.GREY_BD, sw=2))
    b.append(S.ui("plug-connected", tx + 16, top + 16, 30, S.NAVY))
    b.append(S.text(tx + 58, top + 42, "Target services", 20, S.NAVY, 700))
    chips = [("si-openai", "OpenAI", True), ("si-googlegemini", "Gemini", True),
             ("github-original", "GitHub", True), ("jira-original", "Jira", True),
             ("building", "Your API", False), ("server-2", "MCP tools", False)]
    cw, chh = 134, 44
    for i, (ic, lab, brand) in enumerate(chips):
        cx = tx + 12 + (i % 2) * (cw + 8)
        cy = top + 60 + (i // 2) * (chh + 8)
        b.append(S.chip(cx, cy, cw, chh, ic, lab, brand=brand, size=17))

    # arrows: 1 request, 3 authorized, 4 credential back, 5 scoped call
    b.append(S.arrow(232, 120, 326, 120))
    b.append(arrow_label(279, 120, ["Request +", "identity"], 1))
    b.append(S.arrow(562, 105, 656, 105))
    b.append(arrow_label(609, 105, ["Allowed", "request"], 3))
    b.append(S.arrow(658, 185, 564, 185, color=S.SECOND))
    b.append(arrow_label(611, 185, ["Short-lived", "credential"], 4, above=False))
    b.append(S.arrow(892, 135, 976, 135))
    b.append(arrow_label(934, 135, ["Scoped", "call"], 5))

    # attack strip and outcome strip
    sy, sh = 248, 80
    b.append(S.rect(0, sy, 700, sh, fill=S.RED_LT, stroke=S.RED, sw=2, dash="9 6"))
    b.append(S.ui("skull", 18, sy + 20, 40, S.RED))
    b.append(S.text(74, sy + 34, "If an attacker hijacks the agent", 19, S.RED, 700))
    b.append(S.text(74, sy + 60, "through prompt injection, poisoned tool output or a rogue MCP server",
                    17, S.INK, 500))
    b.append(S.rect(720, sy, 560, sh, fill=S.GREEN_LT, stroke=S.GREEN, sw=2))
    b.append(S.ui("shield-lock", 738, sy + 20, 40, S.GREEN_DK))
    b.append(S.text(794, sy + 34, "It finds no key to steal", 19, S.GREEN_DK, 700))
    b.append(S.text(794, sy + 60, "Secrets stay in the vault, scoped and brief", 17, S.INK, 500))
    return S.render("agent_vault_flow", 1280, 334, "".join(b), DIAG)


# =================================================================== slide
def build_slide(flow_png):
    shutil.copy(SRC, OUT)
    prs = Presentation(OUT)
    sl = prs.slides._sldIdLst
    for sid in list(sl):
        try: prs.part.drop_rel(sid.get(qn("r:id")))
        except Exception: pass
        sl.remove(sid)

    s = prs.slides.add_slide(prs.slide_layouts[4])
    set_title(s, "AI Agents Use Credentials They Never Hold")
    blank_footer(s)

    bottom = diagram(s, flow_png, CX, 0.88, CW)

    ly = bottom + 0.08
    label(s, CX, ly, CW, "SIX CONTROLS AROUND THE VAULT", size=9.5)
    controls = [("ban", "Never hardcode\nsecrets"),
                ("clock-hour-4", "Use short-lived\ntokens"),
                ("lock-access", "Grant least\nprivilege"),
                ("world", "Check outbound\ncalls"),
                ("report-search", "Audit every\nsecret access"),
                ("shield-x", "Distrust tool\noutput")]
    gap = 0.12
    w = (CW - gap * 5) / 6
    cy, ch = ly + 0.28, 1.06
    for i, (ic, lab) in enumerate(controls):
        x = CX + i * (w + gap)
        box(s, x, cy, w, ch, fill=WHITE, line=GREY_BD, line_w=1.0,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        icon_tile(s, ic, x + w / 2 - 0.2, cy + 0.1, 0.40, bg=RED if i == 5 else PRIMARY)
        box(s, x + 0.04, cy + 0.54, w - 0.08, 0.46, text=lab, size=10.5, color=NAVY,
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ml=0, mr=0)

    callout(s, CX, cy + ch + 0.08, CW, 0.40, "shield-lock",
            "A vault alone is not enough. Guard identity, access, retrieval and outbound traffic before agents ship.",
            size=11.5)

    notes(s, """
Talk track. AI agents need credentials to call LLMs, GitHub, Jira and internal APIs. The safe pattern lets them use those credentials without ever seeing them.

The flow, in five steps.
1. The agent sends a request with its identity and intent. Example: "I need to call the GitHub API to fetch repo details."
2. The agent gateway verifies the identity, checks permissions, reviews the requested action and blocks suspicious requests.
3. The gateway sends an authorized request for the secret to the credential vault.
4. The vault injects a short-lived credential, only when needed. The vault holds API keys, OAuth tokens and other secrets, and rotates them.
5. The call goes to the target service with limited scope: GPT and other LLMs, Gemini, enterprise APIs, GitHub, Jira, MCP servers and internal MCP.

What the agent never has: direct access to secrets, hardcoded keys, or long-lived tokens.

If the agent is compromised through prompt injection, malicious tool output or an insecure MCP server, exposed credentials would be a gateway to sensitive systems. With the vault and these controls, there is no credential in the agent to steal.

Six controls.
1. Never hardcode secrets. Keep credentials out of prompts, memory, config files and code.
2. Use short-lived credentials. Rotate tokens automatically.
3. Apply least privilege. Restrict each agent to specific APIs, resources and operations.
4. Validate outbound requests. Stop credentials from going to untrusted domains.
5. Audit every secret access. Log which agent asked, what resource was accessed, and whether it was authorized.
6. Defend against prompt injection. Treat retrieved documents and tool responses as untrusted input, never as authorization instructions.

Takeaway. A credential vault alone is not enough. Build a security boundary around identity, authorization, secret retrieval and outbound network access before AI agents reach production.
""")
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    png = build_flow()
    print("diagram", png)
    print("deck", build_slide(png))
