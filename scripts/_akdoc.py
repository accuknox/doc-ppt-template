# -*- coding: utf-8 -*-
"""
Shared AccuKnox Word letterhead helpers.

Every document starts from WORD_TEMPLATE_ACCUKNOX.docx, so the logo header, the brand
fonts and the heading styles come from the template. These helpers add the pieces a
client report needs on top of it: a header title, a letterhead footer, a document
control block, stat tiles, severity badges, callouts and a signature block.

Palette follows README.md and scripts/_akdeck.py. Convert to PDF with to_pdf(), which
drives Word through PowerShell COM, because Word renders Space Grotesk as shipped.
"""
import os
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "WORD_TEMPLATE_ACCUKNOX.docx"

FONT = "Space Grotesk"
MONO = "Consolas"

NAVY = RGBColor(0x11, 0x20, 0x6D)
PRIMARY = RGBColor(0x00, 0x46, 0xFF)
SECOND = RGBColor(0x64, 0x64, 0xFF)
PURPLE = RGBColor(0x4D, 0x4D, 0xD9)
RED = RGBColor(0xC8, 0x00, 0x19)
GREEN = RGBColor(0x16, 0xA5, 0x5C)
GREEN_DK = RGBColor(0x0B, 0x7A, 0x42)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1B, 0x22, 0x3B)
MUTE = RGBColor(0x5A, 0x63, 0x7D)

HEX_NAVY = "11206D"
HEX_PRIMARY = "0046FF"
HEX_GREY_BG = "EEF0F6"
HEX_GREY_BD = "C4CCDE"
HEX_ZEBRA = "F7F8FB"
HEX_RED_LT = "FBECEE"
HEX_LAV = "ECECFB"
HEX_GREEN_LT = "E7F6EE"

# Severity badge tints, per CLAUDE.md rule 5.
SEV = {
    "critical": (HEX_RED_LT, RED),
    "high": (HEX_LAV, PURPLE),
    "medium": (HEX_GREY_BG, MUTE),
    "low": (HEX_GREEN_LT, GREEN_DK),
    "informational": ("FFFFFF", MUTE),
    "clean": (HEX_GREEN_LT, GREEN_DK),
}


# ------------------------------------------------------------------ low level
def _font(run, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


def style_run(run, size=None, bold=None, italic=None, color=None, font=FONT):
    _font(run, font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_text(p, text, **kw):
    """Add text to a paragraph. A newline becomes a real line break, never a raw \\n."""
    parts = str(text).split("\n")
    run = None
    for i, part in enumerate(parts):
        run = p.add_run(part)
        style_run(run, **kw)
        if i < len(parts) - 1:
            run.add_break(WD_BREAK.LINE)
    return run


def shade(cell, hexfill):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:shd")):
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcpr.append(shd)


def table_borders(table, color=HEX_GREY_BD, sz=4, val="single", inside=True):
    tblpr = table._tbl.tblPr
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    el = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right") + (("insideH", "insideV") if inside else ())
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        if edge in edges:
            e.set(qn("w:val"), val)
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        el.append(e)
    tblpr.append(el)


def cell_borders(cell, color="FFFFFF", sz=18):
    tcpr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        el.append(e)
    tcpr.append(el)


def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._tbl.tblPr
    el = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        el.append(e)
    tblpr.append(el)


def _fixed_layout(table, widths):
    tblpr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        if i < len(widths):
            gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in table.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)


def _field(p, instr, **kw):
    r1 = p.add_run()
    style_run(r1, **kw)
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r1._r.append(b)
    r2 = p.add_run()
    style_run(r2, **kw)
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r2._r.append(t)
    r3 = p.add_run()
    style_run(r3, **kw)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = p.add_run("1")
    style_run(r4, **kw)
    r5 = p.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)


# ------------------------------------------------------------------ document
class AKDoc:
    """One branded document built on the AccuKnox Word template."""

    def __init__(self, header_title, classification="Confidential"):
        self.doc = Document(TEMPLATE)
        self._set_header(header_title)
        self._set_footer(classification)
        body = self.doc.element.body
        for el in list(body):
            if el.tag != qn("w:sectPr"):
                body.remove(el)

    # ---- letterhead
    def _set_header(self, text):
        p = self.doc.sections[0].header.paragraphs[0]
        for run in p.runs:
            if "Document Title" in run.text:
                run.text = text
                style_run(run, size=11, color=NAVY)

    def _set_footer(self, classification):
        footer = self.doc.sections[0].footer
        for p in list(footer.paragraphs):
            p._p.getparent().remove(p._p)
        p = footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        pbdr = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single")
        top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "6")
        top.set(qn("w:color"), HEX_GREY_BD)
        pbdr.append(top)
        p._p.get_or_add_pPr().append(pbdr)
        add_text(p, "AccuKnox, Inc.   |   www.accuknox.com   |   support@accuknox.com",
                 size=8, color=MUTE)
        p2 = footer.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p2, classification + "   |   Page ", size=8, color=MUTE)
        _field(p2, "PAGE", size=8, color=MUTE)
        add_text(p2, " of ", size=8, color=MUTE)
        _field(p2, "NUMPAGES", size=8, color=MUTE)

    # ---- text blocks
    def p(self, text="", size=10, color=INK, bold=False, italic=False, align=None,
          before=0, after=6, font=FONT, keep=False, container=None):
        target = container if container is not None else self.doc
        para = target.add_paragraph(style="normal")
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(after)
        para.paragraph_format.line_spacing = 1.12
        if align is not None:
            para.alignment = align
        if keep:
            para.paragraph_format.keep_with_next = True
        if text:
            add_text(para, text, size=size, color=color, bold=bold, italic=italic, font=font)
        return para

    def rich(self, pieces, size=10, after=6, before=0, container=None, align=None):
        """pieces: list of (text, {style}) tuples on one paragraph."""
        para = self.p("", after=after, before=before, container=container, align=align)
        for text, kw in pieces:
            opts = dict(size=size, color=INK)
            opts.update(kw)
            add_text(para, text, **opts)
        return para

    def eyebrow(self, text, before=6):
        return self.p(text.upper(), size=9, bold=True, color=PRIMARY, before=before, after=2,
                      keep=True)

    def title(self, text, subtitle=None):
        para = self.p(text, size=26, bold=True, color=PRIMARY, after=4, keep=True)
        para.paragraph_format.line_spacing = 1.0
        if subtitle:
            self.p(subtitle, size=12, color=MUTE, after=14)
        return para

    def h1(self, text, before=16):
        para = self.doc.add_paragraph(style="Heading 1")
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.keep_with_next = True
        add_text(para, text, size=17, bold=True, color=PRIMARY)
        return para

    def h2(self, text, before=12):
        para = self.doc.add_paragraph(style="Heading 2")
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.keep_with_next = True
        add_text(para, text, size=13, bold=True, color=NAVY)
        return para

    def h3(self, text, before=10):
        para = self.doc.add_paragraph(style="Heading 3")
        para.paragraph_format.space_before = Pt(before)
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.keep_with_next = True
        add_text(para, text, size=11, bold=True, color=NAVY)
        return para

    def bullets(self, items, size=10, container=None):
        for item in items:
            para = self.p("", after=3, container=container)
            para.paragraph_format.left_indent = Inches(0.25)
            para.paragraph_format.first_line_indent = Inches(-0.18)
            add_text(para, "•  ", size=size, color=PRIMARY, bold=True)
            if isinstance(item, tuple):
                add_text(para, item[0] + " ", size=size, color=INK, bold=True)
                add_text(para, item[1], size=size, color=INK)
            else:
                add_text(para, item, size=size, color=INK)

    def page_break(self):
        para = self.doc.add_paragraph(style="normal")
        para.paragraph_format.space_after = Pt(0)
        para.add_run().add_break(WD_BREAK.PAGE)

    def spacer(self, pts=6):
        para = self.p("", after=0)
        para.paragraph_format.space_before = Pt(pts)
        return para

    # ---- tables
    def _new_table(self, rows, cols, widths, container=None):
        target = container if container is not None else self.doc
        t = target.add_table(rows=rows, cols=cols)
        t.style = self.doc.styles["TableNormal"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _fixed_layout(t, widths)
        cell_margins(t)
        return t

    def cell(self, cell, text, size=9, bold=False, color=INK, align=None, font=FONT,
             fill=None, valign=True):
        cell.text = ""
        para = cell.paragraphs[0]
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(1)
        para.paragraph_format.line_spacing = 1.05
        if align is not None:
            para.alignment = align
        add_text(para, text, size=size, bold=bold, color=color, font=font)
        if fill:
            shade(cell, fill)
        if valign:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return cell

    def badge(self, cell, severity, size=8.5):
        fill, color = SEV[severity.lower()]
        self.cell(cell, severity.capitalize() if severity.lower() != "clean" else "Clean",
                  size=size, bold=True, color=color, fill=fill,
                  align=WD_ALIGN_PARAGRAPH.CENTER)

    def table(self, headers, rows, widths, size=9, head_size=8.5, zebra=True,
              align=None, sev_col=None, bold_first=False, total_row=False, container=None):
        """headers: list[str]; rows: list[list[str]]; align: list of alignments per column."""
        t = self._new_table(len(rows) + 1, len(headers), widths, container=container)
        table_borders(t)
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            self.cell(c, h, size=head_size, bold=True, color=WHITE, fill=HEX_NAVY,
                      align=(align[i] if align else None))
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
        for r, row in enumerate(rows, start=1):
            last = total_row and r == len(rows)
            for i, val in enumerate(row):
                c = t.rows[r].cells[i]
                if sev_col is not None and i == sev_col and not last:
                    self.badge(c, val)
                    continue
                fill = HEX_GREY_BG if last else (HEX_ZEBRA if zebra and r % 2 == 0 else None)
                self.cell(c, val, size=size, bold=(last or (bold_first and i == 0)),
                          color=(NAVY if (bold_first and i == 0) or last else INK),
                          align=(align[i] if align else None), fill=fill)
            tr = t.rows[r]._tr.get_or_add_trPr()
            tr.append(OxmlElement("w:cantSplit"))
        if len(rows) <= 12:
            # Keep a short table on one page: every row but the last keeps with the next.
            for row in t.rows[:-1]:
                for c in row.cells:
                    for para in c.paragraphs:
                        para.paragraph_format.keep_with_next = True
        self.p("", after=4, container=container)
        return t

    def kv_table(self, pairs, widths=(1.9, 4.6), size=9.5, container=None):
        t = self._new_table(len(pairs), 2, widths, container=container)
        table_borders(t)
        for i, (k, v) in enumerate(pairs):
            self.cell(t.rows[i].cells[0], k, size=size - 0.5, bold=True, color=NAVY,
                      fill=HEX_GREY_BG)
            self.cell(t.rows[i].cells[1], v, size=size)
            t.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if i < len(pairs) - 1 and len(pairs) <= 7:
                for c in t.rows[i].cells:
                    for para in c.paragraphs:
                        para.paragraph_format.keep_with_next = True
        self.p("", after=4, container=container)
        return t

    def stats(self, items, width=6.5, container=None):
        """items: list of (number, label, color). One tile per item."""
        n = len(items)
        t = self._new_table(1, n, [width / n] * n, container=container)
        table_borders(t, color="FFFFFF", sz=0, val="nil")
        for i, (num, label, color) in enumerate(items):
            c = t.rows[0].cells[i]
            shade(c, HEX_GREY_BG)
            cell_borders(c, "FFFFFF", 24)
            c.text = ""
            p1 = c.paragraphs[0]
            p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p1.paragraph_format.space_before = Pt(6)
            p1.paragraph_format.space_after = Pt(0)
            add_text(p1, num, size=20, bold=True, color=color)
            p2 = c.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p2.paragraph_format.space_after = Pt(6)
            add_text(p2, label, size=8.5, color=MUTE)
        self.p("", after=4, container=container)
        return t

    def callout(self, heading, body, fill=HEX_LAV, accent=HEX_PRIMARY, color=NAVY):
        t = self._new_table(1, 1, [6.5])
        table_borders(t, color="FFFFFF", sz=0, val="nil")
        c = t.rows[0].cells[0]
        shade(c, fill)
        tcpr = c._tc.get_or_add_tcPr()
        el = OxmlElement("w:tcBorders")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "36")
        left.set(qn("w:space"), "0")
        left.set(qn("w:color"), accent)
        el.append(left)
        tcpr.append(el)
        c.text = ""
        p1 = c.paragraphs[0]
        p1.paragraph_format.space_before = Pt(5)
        p1.paragraph_format.space_after = Pt(2)
        add_text(p1, heading, size=10, bold=True, color=color)
        for line in body if isinstance(body, list) else [body]:
            p2 = c.add_paragraph()
            p2.paragraph_format.space_after = Pt(4)
            p2.paragraph_format.line_spacing = 1.1
            add_text(p2, line, size=9.5, color=INK)
        self.p("", after=4)
        return t

    def signature(self, lines=("Authorized signatory", "Name", "Designation", "Date"),
                  org="For AccuKnox, Inc.", container=None):
        self.p(org, size=10, bold=True, color=NAVY, before=10, after=10, container=container,
               keep=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        t = self._new_table(len(lines), 2, [1.5, 3.2], container=container)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        table_borders(t, color="FFFFFF", sz=0, val="nil")
        for i, label in enumerate(lines):
            self.cell(t.rows[i].cells[0], label, size=9, color=MUTE)
            c = t.rows[i].cells[1]
            self.cell(c, "", size=9)
            tcpr = c._tc.get_or_add_tcPr()
            el = OxmlElement("w:tcBorders")
            b = OxmlElement("w:bottom")
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "6")
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), HEX_GREY_BD)
            el.append(b)
            tcpr.append(el)
            t.rows[i].height = Inches(0.36)
        return t

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        return path


def to_pdf(docx_path, pdf_path):
    """Export with Word COM so Space Grotesk and the header logo render as shipped."""
    docx_path, pdf_path = str(Path(docx_path).resolve()), str(Path(pdf_path).resolve())
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        "$w.DisplayAlerts = 0; "
        f"$d = $w.Documents.Open('{docx_path}', $false, $true); "
        "$d.Fields.Update() | Out-Null; "
        f"$d.SaveAs2('{pdf_path}', 17); $d.Close($false); $w.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    return pdf_path
