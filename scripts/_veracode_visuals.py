# -*- coding: utf-8 -*-
"""
Visual slides for the four Veracode product decks: the problem hero, the solution hub,
the stack-ranking scorecard and the link pills. Imported by _veracode_kit.py.
"""
import base64, os
import _aksvg as S
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from _akdeck import *  # noqa: F401,F403

HERE = os.path.dirname(os.path.abspath(__file__))
EMBLEM = os.path.join(HERE, "..", "assets", "logos", "accuknox-emblem.png")


def _png_uri(path):
    return "data:image/png;base64," + base64.b64encode(open(path, "rb").read()).decode()


def _grad(id_, a, b, x2="1", y2="1"):
    return (f'<linearGradient id="{id_}" x1="0" y1="0" x2="{x2}" y2="{y2}">'
            f'<stop offset="0" stop-color="{a}"/><stop offset="1" stop-color="{b}"/></linearGradient>')


# ------------------------------------------------------------------ problem hero
def svg_problem(name, outdir, hero, cards):
    """hero = dict(eyebrow, big, lines, source). cards = 4 x (icon, head, body_lines)."""
    W, H = 1280, 500
    b = [f'<defs>{_grad("hg", "#0A144A", "#0000C8")}{_grad("rg", "#FFF4F5", "#FFFFFF", "1", "0")}</defs>']
    b.append('<rect x="0" y="0" width="430" height="500" rx="22" fill="url(#hg)"/>')
    b.append('<circle cx="400" cy="40" r="130" fill="#6464FF" opacity="0.18"/>')
    b.append('<circle cx="40" cy="470" r="110" fill="#0046FF" opacity="0.35"/>')
    b.append(S.text(36, 62, hero["eyebrow"], 18, "#B8C4E8", 700))
    b.append(S.text(32, 200, hero["big"], hero.get("big_size", 112), "#FFFFFF", 700))
    b.append(S.text(36, 262, hero["lines"], 26, "#FFFFFF", 600, lh=1.25))
    b.append('<rect x="36" y="430" width="70" height="4" rx="2" fill="#C80019"/>')
    b.append(S.text(36, 464, hero["source"], 17, "#B8C4E8", 500))
    for i, (ic, head, body) in enumerate(cards):
        x = 456 + (i % 2) * 418
        y = (i // 2) * 256
        b.append(S.rect(x, y, 404, 240, fill="url(#rg)", stroke="#F2C9CF", sw=1.5, r=18))
        b.append(f'<rect x="{x}" y="{y + 26}" width="6" height="188" rx="3" fill="#C80019"/>')
        b.append(f'<circle cx="{x + 60}" cy="{y + 60}" r="32" fill="#C80019"/>')
        b.append(S.ui(ic, x + 42, y + 42, 36, "#FFFFFF"))
        b.append(S.text(x + 372, y + 74, f"0{i + 1}", 46, "#F2C9CF", 700, anchor="end"))
        b.append(S.text(x + 30, y + 132, head, 25, "#11206D", 700))
        b.append(S.text(x + 30, y + 170, body, 19, "#5A637D", 500, lh=1.3))
    return S.render(name, W, H, "".join(b), outdir)


# ------------------------------------------------------------------ solution hub
def svg_solution(name, outdir, title_lines, cards, logos):
    """cards = 6 x (icon, head, body_lines). Left column takes 0 to 2, right 3 to 5."""
    W, H = 1280, 470
    cx, cy, r = 640, 222, 138
    b = [f'<defs>{_grad("hub", "#0046FF", "#0000C8")}</defs>']
    for rr, op in ((r + 46, 0.10), (r + 24, 0.18)):
        b.append(f'<circle cx="{cx}" cy="{cy}" r="{rr}" fill="#6464FF" opacity="{op}"/>')
    pos = []
    for i in range(6):
        left = i < 3
        x = 8 if left else 868
        y = 6 + (i % 3) * 154
        pos.append((x, y, left))
        tx = x + 404 if left else x
        ty = y + 68
        hx = cx - r * 0.86 if left else cx + r * 0.86
        hy = cy + (ty - cy) * 0.45
        mx = (tx + hx) / 2
        b.append(f'<path d="M{tx},{ty} C{mx},{ty} {mx},{hy} {hx},{hy}" fill="none" stroke="#0046FF" '
                 f'stroke-width="2.5" stroke-dasharray="7 6" opacity="0.7"/>')
        b.append(f'<circle cx="{hx}" cy="{hy}" r="6" fill="#0046FF"/>')
    b.append(f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="url(#hub)"/>')
    b.append(f'<image href="{_png_uri(EMBLEM)}" x="{cx - 34}" y="{cy - 112}" width="68" height="68"/>')
    ty0 = cy - 6 if len(title_lines) == 2 else cy + 6
    b.append(S.text(cx, ty0, title_lines, 30, "#FFFFFF", 700, anchor="middle", lh=1.15))
    b.append(S.text(cx, cy + 92, "by AccuKnox", 17, "#B8C4E8", 600, anchor="middle"))
    lw = len(logos) * 46 - 8
    for j, lg in enumerate(logos):
        b.append(S.rect(cx - lw / 2 + j * 46 - 2, 404, 42, 42, fill="#FFFFFF", stroke="#C4CCDE", sw=1.2, r=9))
        b.append(S.logo(lg, cx - lw / 2 + j * 46 + 3, 409, 32))
    for (x, y, left), (ic, head, body) in zip(pos, cards):
        b.append(S.rect(x, y, 404, 136, fill="#FFFFFF", stroke="#C4CCDE", sw=1.5, r=16))
        b.append(S.rect(x + 18, y + 22, 56, 56, fill="#0046FF", stroke="#0046FF", sw=1, r=14))
        b.append(S.ui(ic, x + 30, y + 34, 32, "#FFFFFF"))
        b.append(S.text(x + 90, y + 46, head, 22, "#11206D", 700))
        b.append(S.text(x + 90, y + 78, body, 17, "#5A637D", 500, lh=1.28))
    return S.render(name, W, H, "".join(b), outdir)


# ------------------------------------------------------------------ links as pills
_AC = {"api": "API", "iac": "IaC", "cicd": "CI/CD", "cspm": "CSPM", "dspm": "DSPM", "cwpp": "CWPP", "kspm": "KSPM",
       "cdr": "CDR", "aws": "AWS", "sarif": "SARIF", "knoxguard": "KnoxGuard", "sec": "Security", "faqs": "FAQ"}


def _title(slug):
    return " ".join(_AC.get(w, w.capitalize()) for w in slug.split("-"))


def auto_label(u):
    """Short readable label for a help-doc or platform URL."""
    path = u.split("://", 1)[-1].rstrip("/")
    host, _, rest = path.partition("/")
    parts = rest.split("/")
    if host.endswith("help.accuknox.com"):
        if parts[0] == "faqs":
            return _title(parts[-1]) + " FAQ"
        if parts[0] == "support-matrix":
            return _title(parts[-1]) + " support matrix"
        return "Help: " + _title(parts[-1])
    if parts[0] == "platform":
        return _title(parts[-1]) + " platform page"
    if parts[0] == "differentiators":
        return "AccuKnox differentiators"
    return _title(parts[-1])


def link_pills(s, urls, y=5.3, x=None):
    x = CX if x is None else x
    box(s, x, y, 1.02, 0.26, text="LEARN MORE", size=9, color=PRIMARY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    cx = x + 1.04
    for u in urls:
        if isinstance(u, tuple):
            lab, u = u
        else:
            lab = auto_label(u)
        w = 0.46 + 0.074 * len(lab)
        if cx + w > CX + CW:
            break
        sp = box(s, cx, y, w, 0.26, fill=LAV, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5,
                 anchor=MSO_ANCHOR.MIDDLE, ml=0.28, mr=0.06, mt=0, mb=0, wrap=False)
        p = sp.text_frame.paragraphs[0]
        r = p.add_run(); r.text = lab
        r.font.name = FONT; r.font.size = Pt(9.5); r.font.color.rgb = PRIMARY; r.font.bold = True
        r.hyperlink.address = u
        sp.click_action.hyperlink.address = u
        icon_img(s, "external-link", cx + 0.08, y + 0.045, 0.17)
        cx += w + 0.1


# ------------------------------------------------------------------ scorecard
STATUS = {"yes": ("circle-check", "#16A55C"), "partial": ("circle-half-2", "#C27A00"), "no": ("circle-x", "#C80019")}


def scorecard(s, vendors, rows, y0=1.2, crit_w=2.55):
    """vendors = [name, ...] with AccuKnox first. rows = [(criterion, [(status, note), ...]), ...]."""
    n = len(vendors)
    vw = (CW - crit_w) / n
    hh, rh = 0.44, 0.48
    box(s, CX + crit_w - 0.02, y0 - 0.04, vw + 0.04, hh + 0.08 + rh * len(rows), fill=None, line=PRIMARY, line_w=2.25,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    box(s, CX, y0, crit_w - 0.06, hh, text="What the buyer needs", size=11, color=MUTE, bold=True,
        anchor=MSO_ANCHOR.MIDDLE, ml=0)
    for j, v in enumerate(vendors):
        x = CX + crit_w + j * vw
        box(s, x + 0.04, y0 + 0.02, vw - 0.08, hh - 0.04, text=v, size=12 if j == 0 else 11, color=WHITE, bold=True,
            fill=PRIMARY if j == 0 else NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2,
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    score = [0] * n
    for i, (crit, cells) in enumerate(rows):
        y = y0 + hh + 0.06 + i * rh
        if i % 2 == 0:
            box(s, CX, y, CW, rh - 0.04, fill=GREY_BG, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        box(s, CX + 0.1, y, crit_w - 0.12, rh - 0.04, text=crit, size=11, color=NAVY, bold=True,
            anchor=MSO_ANCHOR.MIDDLE, ml=0)
        for j, (st, note) in enumerate(cells):
            ic, col = STATUS[st]
            if st == "yes":
                score[j] += 1
            x = CX + crit_w + j * vw
            icon_img(s, ic, x + 0.1, y + (rh - 0.04) / 2 - 0.15, 0.3, color=col)
            box(s, x + 0.44, y, vw - 0.48, rh - 0.04, text=note, size=9.5, color=INK if st == "yes" else MUTE,
                bold=(j == 0), anchor=MSO_ANCHOR.MIDDLE, ml=0, mr=0)
    return score
