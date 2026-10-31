# -*- coding: utf-8 -*-
"""
AccuKnox vs HashiCorp Vault comparison, on the AccuKnox Word letterhead.

Source of truth is the markdown page in the help repo, so the Word file and the
web page never drift:
  D:/AccuKnox/help/references/drafts/comparison/accuknox-vs-hashicorp-vault.md

Run:  py -3.11 scripts/build_accuknox_vs_vault.py
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from _akdoc import (cell_borders, AKDoc, to_pdf, add_text, style_run, shade, table_borders, MONO, NAVY, INK,
                    MUTE, PRIMARY, RED, GREEN, GREEN_DK, PURPLE, HEX_GREY_BG, HEX_GREY_BD,
                    HEX_GREEN_LT, HEX_LAV, HEX_NAVY, HEX_ZEBRA)
import _aksvg as S

ROOT = Path(__file__).resolve().parents[1]
MD = Path(r"D:/AccuKnox/help/references/drafts/comparison/accuknox-vs-hashicorp-vault.md")
OUTD = Path(r"D:/AccuKnox/help/references/drafts/comparison")
DIAG = ROOT / "work" / "vault_compare" / "diagrams"
NAME = "AccuKnox_vs_HashiCorp_Vault"

LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")
CONF_RE = re.compile(r"\[confirm[^\]]*\]")
TOKEN_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)|(\[confirm[^\]]*\])|`([^`]+)`")


# ------------------------------------------------------------------ markdown
def read_md():
    return MD.read_text(encoding="utf-8").split("---", 2)[2]


def table_after(md, heading):
    """Return the rows of the first markdown table after a heading line."""
    start = md.index(heading)
    lines = md[start:].splitlines()
    rows, seen = [], False
    for ln in lines[1:]:
        if ln.startswith("|"):
            seen = True
            cells = [c.strip() for c in ln.strip().strip("|").split(" | ")]
            if set(ln.replace("|", "").replace(" ", "")) <= {"-"}:
                continue
            rows.append(cells)
        elif seen:
            break
    return rows[0], rows[1:]


def section(md, heading):
    start = md.index(heading) + len(heading)
    nxt = md.find("\n#", start)
    return md[start:nxt if nxt != -1 else None].strip()


# ------------------------------------------------------------------ word bits
def add_link(par, label, url, size=8, color=PRIMARY, bold=False):
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    h.set(qn("w:history"), "1")
    run = par.add_run(label)
    style_run(run, size=size, color=color, bold=bold)
    run.font.underline = True
    h.append(run._r)
    par._p.append(h)


def rich_text(par, text, size=9.5, color=INK, bold=False, link_size=None):
    """Write text with [label](url) links, `code` and red [confirm ...] brackets."""
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            add_text(par, text[pos:m.start()], size=size, color=color, bold=bold)
        if m.group(1):
            add_link(par, m.group(1), m.group(2), size=link_size or size, bold=bold)
        elif m.group(3):
            add_text(par, m.group(3), size=size, color=RED, bold=True)
        else:
            add_text(par, m.group(4), size=size - 0.5, color=NAVY, font=MONO)
        pos = m.end()
    if pos < len(text):
        add_text(par, text[pos:], size=size, color=color, bold=bold)


def body(d, text, size=10, after=6):
    par = d.p("", after=after)
    rich_text(par, text, size=size)
    return par


def split_sources(cell):
    """Split a table cell into (plain text, [(label, url)])."""
    links = LINK_RE.findall(cell)
    plain = LINK_RE.sub("", cell)
    plain = re.sub(r"\s*,\s*(?=\s*$)", "", plain)
    plain = re.sub(r"\s{2,}", " ", plain).strip().rstrip(",").strip()
    return plain, links


def result_style(label):
    lab = label.lower()
    if lab.startswith("accuknox"):
        return HEX_GREEN_LT, GREEN_DK
    if lab.startswith("vault"):
        return HEX_GREY_BG, MUTE
    return HEX_LAV, PURPLE


def fill_cell(d, cell, text, size=8.5):
    plain, links = split_sources(text)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.05
    rich_text(p1, plain, size=size)
    if links:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.0
        add_text(p2, "Source: ", size=7, color=MUTE, bold=True)
        for i, (lab, url) in enumerate(links):
            if i:
                add_text(p2, "  |  ", size=7, color=MUTE)
            add_link(p2, lab, url, size=7)


def matrix(d, rows):
    widths = [1.15, 2.2, 2.2, 0.95]
    t = d._new_table(len(rows) + 1, 4, widths)
    table_borders(t)
    heads = ["Parameter", "AccuKnox Secrets Manager", "HashiCorp Vault", "Result"]
    for i, h in enumerate(heads):
        d.cell(t.rows[0].cells[i], h, size=8.5, bold=True, color=__import__("docx").shared.RGBColor(255, 255, 255),
               fill=HEX_NAVY, align=WD_ALIGN_PARAGRAPH.CENTER if i == 3 else None)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r, (param, ak, hv, res, take) in enumerate(rows, start=1):
        row = t.rows[r]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        zebra = HEX_ZEBRA if r % 2 == 0 else None
        # parameter + takeaway
        c0 = row.cells[0]
        d.cell(c0, param, size=9, bold=True, color=NAVY, fill=HEX_GREY_BG, valign=False)
        p = c0.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        add_text(p, take, size=7.5, color=MUTE, italic=True)
        # competitor columns
        for ci, txt in ((1, ak), (2, hv)):
            c = row.cells[ci]
            fill_cell(d, c, txt)
            if zebra:
                shade(c, zebra)
        fill, col = result_style(res)
        d.cell(row.cells[3], res, size=8, bold=True, color=col, fill=fill,
               align=WD_ALIGN_PARAGRAPH.CENTER)
        cell_borders(row.cells[3], "FFFFFF", 12)
    d.p("", after=4)
    return t


def steps_table(d, rows):
    widths = [0.6, 5.9]
    t = d._new_table(len(rows) + 1, 2, widths)
    table_borders(t)
    W = __import__("docx").shared.RGBColor(255, 255, 255)
    d.cell(t.rows[0].cells[0], "Step", size=8.5, bold=True, color=W, fill=HEX_NAVY,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    d.cell(t.rows[0].cells[1], "What to do", size=8.5, bold=True, color=W, fill=HEX_NAVY)
    for r, (n, what) in enumerate(rows, start=1):
        t.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        d.cell(t.rows[r].cells[0], n, size=11, bold=True, color=PRIMARY, fill=HEX_LAV,
               align=WD_ALIGN_PARAGRAPH.CENTER)
        c = t.rows[r].cells[1]
        c.text = ""
        par = c.paragraphs[0]
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(2)
        rich_text(par, what, size=9.5)
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if r % 2 == 0:
            shade(c, HEX_ZEBRA)
    d.p("", after=4)


# ------------------------------------------------------------------ diagrams
def diagram_pick():
    W, H = 960, 196
    cards = [
        (0, "api", "Same API", ["Keep Vault API, CLI", "and app code"], PRIMARY_S()),
        (245, "users-group", "Namespaces", ["One per team on a", "dedicated instance"], S.PRIMARY),
        (490, "report-search", "Audit trail", ["Logs flow straight to", "AccuKnox SIEM and CDR"], S.PRIMARY),
        (735, "shield-lock", "Hardened", ["CWPP guards every", "Secrets Manager pod"], S.PRIMARY),
    ]
    b = []
    for x, icon, head, sub, acc in cards:
        b.append(S.card(x, 8, 225, 176, icon, head, sub, accent=acc, hsize=23, ssize=17))
    return S.render("pick_accuknox", W, H, "".join(b), str(DIAG))


def PRIMARY_S():
    return S.PRIMARY


def diagram_migration():
    W, H = 960, 236
    steps = [
        ("sitemap", ["Map", "namespaces"]),
        ("license", ["Move ACL", "policies"]),
        ("plug-connected", ["Point apps at", "new endpoint"]),
        ("zoom-check", ["Re-test auth", "and roles"]),
        ("arrows-exchange", ["Cut over", "gradually"]),
        ("report-search", ["Forward audit", "logs to SIEM"]),
    ]
    bw, gap = 140, 24
    b = []
    for i, (icon, label) in enumerate(steps):
        x = i * (bw + gap)
        b.append(S.rect(x, 10, bw, 212, fill=S.WHITE, stroke=S.GREY_BD))
        b.append(S.rect(x, 10, bw, 8, fill=S.PRIMARY, stroke="none", r=4))
        b.append(S.badge(x + 30, 54, i + 1, r=18, fill=S.PRIMARY, size=21))
        b.append(f'<circle cx="{x + bw / 2}" cy="120" r="34" fill="{S.LAV}"/>')
        b.append(S.ui(icon, x + bw / 2 - 22, 98, 44, S.PRIMARY))
        b.append(S.text(x + bw / 2, 184, label, 17, S.NAVY, 700, anchor="middle"))
        if i < len(steps) - 1:
            b.append(S.arrow(x + bw + 3, 116, x + bw + gap - 3, 116, S.PRIMARY, sw=3))
    return S.render("migration_six_steps", W, H, "".join(b), str(DIAG))


def picture(d, path, width=6.5, caption=None):
    par = d.p("", after=3, keep=True)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        d.p(caption, size=8.5, color=MUTE, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)


# ------------------------------------------------------------------ build
def main():
    md = read_md()
    DIAG.mkdir(parents=True, exist_ok=True)
    pick_png = diagram_pick()
    mig_png = diagram_migration()

    _, mrows = table_after(md, "## AccuKnox Leads on")
    rows = [(r[0], r[1], r[2], r[3], r[4]) for r in mrows]
    _, srows = table_after(md, "## Migrate From Vault")

    n = len(rows)
    lead = sum(1 for r in rows if r[3].lower().startswith("accuknox"))
    vlead = sum(1 for r in rows if r[3].lower().startswith("vault"))
    par = n - lead - vlead

    d = AKDoc("AccuKnox vs HashiCorp Vault", classification="Draft for review")

    d.eyebrow("Platform comparison", before=0)
    d.title("AccuKnox vs HashiCorp Vault", "Secrets management platform comparison")
    d.p("AccuKnox Secrets Manager keeps the Vault API and adds namespaces and runtime hardening.",
        size=13, bold=True, color=NAVY, after=10)
    d.stats([(str(n), "Parameters compared", PRIMARY),
             (str(lead), "AccuKnox leads", GREEN),
             (str(par), "Parity", PURPLE)])
    d.eyebrow("Pick AccuKnox Secrets Manager to")
    picture(d, pick_png, 6.5)
    body(d, section(md, "## Secrets Management Platform Comparison").split("\n\n")[1], size=10)

    d.h1(f"AccuKnox Leads on {lead} of {n} Secrets Manager Parameters")
    matrix(d, rows)
    d.h1("Why Customers Choose AccuKnox Over HashiCorp Vault")
    for h, key in (("Better", "### Better"), ("Faster", "### Faster"), ("Cheaper", "### Cheaper")):
        d.h2(h)
        body(d, section(md, key).split("\n### ")[0].split("\n## ")[0].strip(), size=10)

    d.page_break()
    d.h1("Migrate From Vault to AccuKnox in Six Steps", before=0)
    picture(d, mig_png, 6.5)
    steps_table(d, srows)

    # sources
    links, seen = [], set()
    for lab, url in LINK_RE.findall(md):
        if url not in seen:
            seen.add(url)
            links.append((lab, url))
    d.page_break()
    d.h1("Every Claim Traces to a Page the Publisher Wrote", before=0)
    d.p("Pages read on 26 October 2026. AccuKnox claims link to AccuKnox pages. "
        "Vault claims link to HashiCorp pages.", size=9.5, color=MUTE, after=6)
    t = d._new_table(len(links) + 1, 3, [1.2, 1.9, 3.4])
    table_borders(t)
    W = __import__("docx").shared.RGBColor(255, 255, 255)
    for i, h in enumerate(["Publisher", "Page", "Address"]):
        d.cell(t.rows[0].cells[i], h, size=8.5, bold=True, color=W, fill=HEX_NAVY)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for r, (lab, url) in enumerate(sorted(links, key=lambda x: ("hashicorp" in x[1], x[1])), start=1):
        pub = "HashiCorp" if "hashicorp" in url else "AccuKnox"
        t.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        d.cell(t.rows[r].cells[0], pub, size=8.5, bold=True, color=NAVY)
        for ci, txt in ((1, lab), (2, url)):
            c = t.rows[r].cells[ci]
            c.text = ""
            pp = c.paragraphs[0]
            pp.paragraph_format.space_before = Pt(1)
            pp.paragraph_format.space_after = Pt(1)
            add_link(pp, txt, url, size=8 if ci == 1 else 7.5)
        if r % 2 == 0:
            for c in t.rows[r].cells:
                shade(c, HEX_ZEBRA)
    d.p("", after=6)
    d.callout("See it on your stack",
              ["Schedule a demo of AccuKnox Secrets Manager at https://www.accuknox.com/contact-us "
               "or write to support@accuknox.com."])

    out_docx = d.save(OUTD / f"{NAME}.docx")
    pdf = to_pdf(out_docx, OUTD / f"{NAME}.pdf")
    print(out_docx, pdf)


if __name__ == "__main__":
    main()
