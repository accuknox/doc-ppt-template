# -*- coding: utf-8 -*-
"""
AccuKnox Agentic Identity Security, the positioning deck.

Sources:
  1. "Zero Trust Agentic Identity Security" PRD deck v0.1, September 2026 (Google Slides
     1eBO20bBCPIJSZxyVVsXWWHj47N-8RytVo_HTJqykPKE). The primary source. Every capability,
     threat, persona, architecture and differentiator comes from it.
  2. "Agentic AI Identity Security" technical deck, 31 August 2026 (Google Slides
     1zR9wOjCOaFjWNghenHZYjLwN93ypuln9lAXU3YN5oOg). Used lightly: blast radius, caller
     chain example, mesh stack, demo scope, prompt firewall coverage.
  3. accuknox.com/solutions/agentic-ai-security and accuknox.com/platform/agentz,
     scraped with Firecrawl on 2026-10-07, for positioning lines.
  4. help.accuknox.com CIEM overview page (docs/getting-started/ciem-overview.md) and
     its screenshots in the help repo.
  5. Market numbers: CyberArk 2025 Identity Security Landscape, SailPoint 2025 AI agent
     adoption report. Cited in the notes of the slide that uses them.

PRD-internal material (goals, targets, roadmap, open questions) sits on hidden backup slides.

Build:  py -3.11 scripts/agentic_identity_svg.py && py -3.11 scripts/build_agentic_identity.py
"""
import os, shutil, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _akdeck import *          # noqa: F401,F403

google_safe()

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
SRC = os.path.join(ROOT, "PPT Template.pptx")
OUT = os.path.join(ROOT, "output", "AccuKnox_Agentic_Identity_Security.pptx")
DIAG = os.path.join(ROOT, "work", "agentic_identity", "diagrams")
SHOTS = os.path.join(ROOT, "work", "agentic_identity", "shots")
CIEM = r"D:\AccuKnox\help\docs\getting-started\images\ciem"
ORANGE = RGBColor(0xE8, 0xA3, 0x3D)
BR = chr(10)   # runs() turns this into a real <a:br/>

os.makedirs(SHOTS, exist_ok=True)
shutil.copy(SRC, OUT)
prs = Presentation(OUT)
L_DIV, L_STD = prs.slide_layouts[2], prs.slide_layouts[4]
xml_slides = prs.slides._sldIdLst
for sid in list(xml_slides):
    try: prs.part.drop_rel(sid.get(qn('r:id')))
    except Exception: pass
    xml_slides.remove(sid)


# =================================================================== helpers
def std(title, kicker, notes_text):
    s = prs.slides.add_slide(L_STD); set_title(s, title); blank_footer(s)
    eyebrow(s, kicker, y=0.84, w=8.6, size=10.5)
    notes(s, notes_text)
    return s


def divider(part, title, sub, icon, notes_text):
    s = prs.slides.add_slide(L_DIV); blank_footer(s)
    box(s, 0.86, 1.62, 6.6, 0.32, text=part.upper(), size=13, color=SECOND, bold=True)
    box(s, 0.82, 1.98, 6.9, 1.10, text=title, size=30, color=WHITE, bold=True)
    box(s, 0.86, 3.20, 0.12, 0.62, fill=PRIMARY)
    box(s, 1.10, 3.18, 6.4, 0.66, text=sub, size=12.5, color=NAVY_TXT, anchor=MSO_ANCHOR.MIDDLE)
    icon_img(s, icon, 7.75, 1.85, 1.5, color="#6464FF")
    notes(s, notes_text)
    return s


def dg(s, name, y=1.16, w=CW):
    return diagram(s, os.path.join(DIAG, name + ".png"), (10 - w) / 2, y, w)


def take(s, text, icon="bulb", y=4.9, h=0.44):
    callout(s, CX, y, CW, h, icon, text, size=12)


def num_tile(s, x, y, n, size=0.5, fill=NAVY):
    box(s, x, y, size, size, text=n, size=15, color=WHITE, bold=True, fill=fill,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, ml=0, mr=0)


def shot(name, crop=None):
    src = os.path.join(CIEM, name)
    if not crop:
        return src
    out = os.path.join(SHOTS, name.replace(".png", "_crop.png"))
    Image.open(src).crop(crop).save(out)
    return out


def tint_cells(t, rules):
    """rules: {text: (fill, color)} applied to body cells with exactly that text."""
    for ri, row in enumerate(t.rows):
        if ri == 0:
            continue
        for c in row.cells:
            if c.text in rules:
                f, col = rules[c.text]
                c.fill.solid(); c.fill.fore_color.rgb = f
                for p in c.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.CENTER
                    for r in p.runs:
                        r.font.color.rgb = col; r.font.bold = True


SECTIONS = {}


def mark(key):
    SECTIONS.setdefault(key, []).append(len(prs.slides))


# =================================================================== 1. cover
cover_slide(prs, "Agentic Identity Security", tsize=30,
            subtitle="Zero Trust for every AI agent and every action it takes")


# =================================================================== 2. TOC
s = std("Six Parts Take You From the Problem to Proof", "Table of contents", """
Agenda. The deck follows one line of argument.
Part 1, the business problem: AI agents act with real credentials, and identity tools were built for people.
Part 2, why today's options fall short: static keys, legacy IAM, and three vendor categories that each cover only part of the problem.
Part 3, the AccuKnox solution: the AgentZ control plane, the reference architecture, and how one request flows.
Part 4, ten capabilities in depth, two per slide.
Part 5, CIEM: the cloud identity inventory AccuKnox ships for human and machine identities.
Part 6, why AccuKnox: differentiators, the category comparison, outcomes and the demo.
Backup slides after the closing slide are hidden. They hold the PRD goals, proposed targets, roadmap and open questions.
""")
toc = [("alert-triangle", "The Problem", "Agents act, IAM cannot see"),
       ("scale", "Why Tools Fall Short", "Keys, IAM, point tools"),
       ("shield-lock", "The AccuKnox Solution", "Control plane and flow"),
       ("list-details", "Capabilities in Depth", "Ten capabilities, flows, stack"),
       ("fingerprint", "CIEM in the Product", "Cloud identities today"),
       ("rocket", "Why AccuKnox", "Proof, outcomes, demo")]
TOC_RANGES = []
for i, (icn, h, sub) in enumerate(toc):
    x = 0.40 + (i % 3) * 3.1
    y = 1.22 + (i // 3) * 2.0
    box(s, x, y, 2.96, 1.82, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    icon_tile(s, icn, x + 0.18, y + 0.2, 0.74)
    box(s, x + 1.9, y + 0.16, 0.9, 0.6, text=f"0{i+1}", size=28, color=LAV, bold=True, align=PP_ALIGN.RIGHT)
    box(s, x + 0.18, y + 1.0, 2.7, 0.36, text=h, size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x + 0.18, y + 1.34, 2.7, 0.26, text=sub, size=10.5, color=MUTE)
    TOC_RANGES.append(box(s, x + 0.18, y + 1.56, 2.7, 0.24, text="Slides", size=9.5, color=PRIMARY, bold=True))


# =================================================================== PART 1
divider("Part 1 · The business problem", "AI Agents Now Act With Real Credentials",
        "They call tools, move data and chain to other agents. Identity tools still assume a person is at the keyboard.",
        "robot", "Part 1 sets the business context. Agents are now the fastest-growing identity type, and the identity stack cannot tell who an agent acts for.")
mark("p1")

s = std("Machine Identities Already Outnumber People 82 to 1", "Context", """
Talk track: agents are the newest and fastest-growing kind of machine identity, and most programs still govern only humans.
Sources.
82 to 1, and 88 percent: CyberArk 2025 Identity Security Landscape, a survey of 2,600 security decision-makers. Machine identities outnumber humans 82 to 1. 88 percent of organizations define only human identities as privileged users. cyberark.com/resources/identity-security/2025-identity-security-landscape
80 percent and 23 percent: SailPoint research on AI agent adoption, 2025. 80 percent of companies say their AI agents took unintended actions, including access to unauthorized systems (39 percent). 23 percent say agents were tricked into revealing access credentials. sailpoint.com/press-releases/sailpoint-ai-agent-adoption-report
These are third-party survey numbers. Quote them with the source on any external version.
""")
mark("p1")
stats = [("82:1", "machine identities for every human identity", "CyberArk 2025", "robot"),
         ("88%", "of organizations treat only humans as privileged", "CyberArk 2025", "user"),
         ("80%", "saw AI agents take unintended actions", "SailPoint 2025", "alert-triangle"),
         ("23%", "saw agents tricked into revealing credentials", "SailPoint 2025", "key")]
for i, (n, l, src, ic) in enumerate(stats):
    x = 0.40 + i * 2.33
    box(s, x, 1.25, 2.18, 2.75, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    icon_tile(s, ic, x + 0.79, 1.42, 0.6, bg=RED if i >= 2 else PRIMARY)
    box(s, x, 2.12, 2.18, 0.62, text=n, size=32, color=RED if i >= 2 else PRIMARY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x + 0.12, 2.76, 1.94, 0.74, text=l, size=11.5, color=INK, align=PP_ALIGN.CENTER)
    box(s, x, 3.6, 2.18, 0.28, text=src, size=9.5, color=GREY_TX, align=PP_ALIGN.CENTER, italic=True)
take(s, "Agents are the fastest-growing identity type, and the least governed one.", icon="trending-up", y=4.3, h=0.5)

s = std("IAM Was Built for People, Not for Agents", "The problem", """
From the PRD, The Problem.
1. They act, not just answer. Agents call tools, move data and trigger workflows on their own. They often chain to other agents with no human in the loop.
2. Their credentials are static and over-scoped. Long-lived API keys and shared service accounts give agents far more access than any single task needs. Those keys leak into prompts and code.
3. Accountability disappears in the chain. A human asks an agent, that agent asks a sub-agent, and the sub-agent calls an API. Today's logs cannot say who authorized what, or stop it mid-flight.
""")
mark("p1")
probs = [("robot", "They act, not only answer", "Agents call tools, move data and trigger workflows, often with no human in the loop."),
         ("key", "Their keys are static and too broad", "Long-lived API keys and shared accounts give far more access than a task needs."),
         ("route", "Accountability breaks in the chain", "When an agent calls a sub-agent, logs cannot say who authorized the action.")]
for i, (ic, h, b) in enumerate(probs):
    x = 0.40 + i * 3.1
    box(s, x, 1.25, 2.96, 3.3, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(s, x + 0.2, 1.42, 0.5, 0.5, text=f"0{i+1}", size=13, color=WHITE, bold=True, fill=ORANGE, shape=MSO_SHAPE.OVAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ml=0, mr=0)
    icon_img(s, ic, x + 2.1, 1.38, 0.62, color="#11206D")
    box(s, x + 0.2, 2.15, 2.6, 0.8, text=h, size=15, color=NAVY, bold=True)
    box(s, x + 0.2, 2.95, 2.6, 1.4, text=b, size=12, color=MUTE)
take(s, "Identity tools check who signed in. Nobody checks what the agent does next.", icon="eye", y=4.75)

s = std("The Security Gap Sits Between the Agent and Its Credentials", "The hidden security boundary", """
The chain runs Human identity, Agent identity, Workload credentials, MCP tools. The gap sits between the agent and the workload credentials it uses.
Key principle from the PRD: never let the agent say it is acting for User A become the authorization decision.
Always verify the identity-to-agent mapping on your own before you allow access to sensitive data or high-impact actions. Secure the identity behind the agent, not only the model inside it.
""")
mark("p1")
dg(s, "boundary", y=1.25, w=8.6)
take(s, "Secure the identity behind the agent, not only the model inside it.", icon="shield-lock", y=4.55)

s = std("Identity-to-Agent Mapping Is the Missing Layer", "What has to exist", """
The identity-to-agent mapping layer binds a verified identity to an agent, with explicit delegation, credentials and policy context.
Identity sources: human identity (end user, employee, customer), agent identity (AI agent or service identity), service account (application or workload identity), delegated identity (on behalf of, temporary, external).
The layer does identity binding, delegation, policy context and an audit trail.
Authorization and policy: a policy engine checks if the mapped identity and agent can perform an action on a resource. RBAC, ABAC, least privilege, per-action checks.
Tools, APIs and resources: external tools, APIs, databases, applications and services. Tool allowlists, resource ACLs, tenant isolation.
Delegation and credentials: short-lived tokens, workload credentials and explicit delegation. OAuth and OIDC, scoped credentials, no implicit inheritance.
Security outcomes: correct identity binding, minimal permissions, controlled delegation, full auditability, tenant isolation.
""")
mark("p1")
dg(s, "mapping", y=1.16, w=8.9)
take(s, "Verified identity plus scoped delegation plus least privilege equals trusted AI actions.", icon="check", y=4.85)

s = std("Five Attack Surfaces Open Up in the Mapping Gap", "Attack surfaces", """
From the PRD, the 5 core attack surfaces of the mapping gap.
Identity confusion: the agent becomes associated with the wrong user. Impact: identity spoofing and unauthorized actions.
Privilege inheritance: the agent receives excessive permissions through broad service accounts. Impact: super NHIs with admin privileges.
Delegation abuse: legitimate delegated credentials are misused. Impact: cascading compromise across trusted relationships.
Cross-tenant leakage: one user's context bleeds into another user's agent. Impact: data exposure across isolated tenants.
Stale authorization: revoked permissions remain usable, the digital fossils. Impact: persistent backdoors through unmanaged API keys.
""")
mark("p1")
rows = [["Attack surface", "What happens", "Business impact"],
        ["Identity confusion", "The agent is tied to the wrong user", "Spoofed identity, unauthorized actions"],
        ["Privilege inheritance", "Broad service accounts hand out too much access", "Super NHIs with admin rights"],
        ["Delegation abuse", "Valid delegated credentials get misused", "Compromise spreads across trust links"],
        ["Cross-tenant leakage", "One user's context reaches another user's agent", "Data exposed across tenants"],
        ["Stale authorization", "Revoked permissions still work", "Backdoors through forgotten API keys"]]
t = table(s, CX, 1.22, CW, 3.3, rows, colw=[2.3, 3.5, 3.4], hsize=12, bsize=11)
for r in range(1, 6):
    t.cell(r, 2).fill.solid(); t.cell(r, 2).fill.fore_color.rgb = RED_LT
take(s, "Each surface turns one compromised agent into access across the business.", icon="alert-triangle", y=4.75)

s = std("Six Threats Target AI Agents Today", "Threat model", """
From the PRD threat model. What we defend against.
Prompt and indirect injection: hostile instructions in inputs, documents or tool output steer the agent.
Confused deputy: an agent uses its own broad privileges for a user who lacks them.
Credential theft and leakage: keys exposed in context windows, logs, repos or memory.
Shadow and orphaned agents: unregistered agents, or agents that outlive their owner or purpose.
Tool and MCP poisoning: malicious or silently changed tool descriptions and servers.
Sandbox escape and lateral movement: agents break out of their runtime to reach hosts, networks and data.
Reference frameworks: OWASP Top 10 for LLM Applications, OWASP agentic AI threat guidance, NIST AI RMF.
""")
mark("p1")
threats = [("message-exclamation", "Prompt injection", "Hostile text in inputs or tool output steers the agent"),
           ("user-question", "Confused deputy", "The agent uses its broad rights for a user who lacks them"),
           ("key", "Credential theft", "Keys leak into context, logs, repos or memory"),
           ("ghost", "Shadow agents", "Unregistered agents, or agents that outlive their owner"),
           ("plug-connected", "Tool and MCP poisoning", "Malicious or silently changed tools and servers"),
           ("arrows-split", "Sandbox escape", "Agents break out to reach hosts, networks and data")]
for i, (ic, h, b) in enumerate(threats):
    x = 0.40 + (i % 3) * 3.1
    y = 1.22 + (i // 3) * 1.55
    icon_card(s, x, y, 2.96, 1.4, ic, h, b, accent=RED, hsize=13, bsize=11)
take(s, "Mapped to OWASP Top 10 for LLM Applications and NIST AI RMF.", icon="checklist", y=4.4)

s = std("Five Teams Need Answers About Every Agent", "Who asks", """
From the PRD personas.
CISO: needs proof that every agent is known, least-privileged and stoppable, in board and regulator language.
IAM or security architect: extends Zero Trust policy from humans and workloads to agents without a parallel stack.
AI platform engineer: ships agents fast and wants identity and credentials handled by default, not bolted on.
SOC analyst: investigates agent incidents end to end and kills a rogue chain in seconds.
GRC and compliance: maps agent controls to the EU AI Act, NIST AI RMF, ISO 42001 and SOC 2.
""")
mark("p1")
pers = [("briefcase", "CISO", "Proof that every agent is known and stoppable"),
        ("id-badge-2", "IAM architect", "One Zero Trust policy for people and agents"),
        ("code", "AI platform engineer", "Identity handled by default, not bolted on"),
        ("radar", "SOC analyst", "Full agent incident story, kill a chain fast"),
        ("gavel", "GRC and compliance", "Controls mapped to EU AI Act, ISO 42001")]
for i, (ic, h, b) in enumerate(pers):
    x = 0.40 + i * 1.86
    box(s, x, 1.25, 1.74, 3.1, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    icon_tile(s, ic, x + 0.53, 1.45, 0.68, bg=PRIMARY if i % 2 == 0 else SECOND)
    box(s, x + 0.08, 2.3, 1.58, 0.6, text=h, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, x + 0.1, 2.95, 1.54, 1.3, text=b, size=11, color=MUTE, align=PP_ALIGN.CENTER)
take(s, "One platform answers all five, from the same identity record.", icon="users", y=4.6)


# =================================================================== PART 2
divider("Part 2 · Current alternatives", "Today's Tools Each Cover Part of the Problem",
        "Static keys, legacy IAM and point products leave the gap between the agent and its actions open.",
        "scale", "Part 2 explains why the options teams use today do not close the identity-to-agent gap.")
mark("p2")

s = std("A Credential's Lifetime Sets Its Blast Radius", "Static keys", """
From the technical deck, the blast radius concept.
A long-lived API key gives months or years of exposure, no identity context, and no way to trace a specific actor.
A single-hop JWT or a SPIFFE SVID lives for minutes, uses eager connect-time authentication and rotates on its own.
Identity must come from what a workload is, proven by attestation. A static secret proves only that someone handed it over.
The circle areas show the contrast, not an exact ratio.
""")
mark("p2")
dg(s, "blast", y=1.16, w=9.2)
take(s, "Identity must come from what a workload is, proven by attestation.", icon="fingerprint", y=4.65)

s = std("Legacy IAM Fails on Credentials, Authorization and Visibility", "The executive decision", """
From the PRD, the executive decision table.
Credentials. Legacy static IAM: long-lived API keys and static service accounts that turn into digital fossils. AccuKnox: ephemeral SPIFFE SVIDs anchored by continuous cryptographic workload attestation.
Authorization. Legacy: blind assumption of user privileges, coarse-grained access. AccuKnox: caller-chain evaluation with OpenFGA, enforcing scope attenuation at every delegation hop.
Visibility. Legacy: boundary blind spots, logs show only the final service account. AccuKnox: immutable RFC 8693 nested lineage, full forensic audit for fast AI-SOC response.
""")
mark("p2")
rows = [["", "Legacy static IAM", "AccuKnox agentic identity"],
        ["Credentials", "Long-lived API keys and static service accounts", "Short-lived SPIFFE SVIDs, proven by workload attestation"],
        ["Authorization", "Assumes the user's rights, coarse access", "Checks the caller chain with OpenFGA, narrows scope per hop"],
        ["Visibility", "Logs show only the final service account", "RFC 8693 lineage from first prompt to final write"]]
t = table(s, CX, 1.22, CW, 2.9, rows, colw=[1.8, 3.5, 3.9], hsize=13, bsize=12)
for r in range(1, 4):
    t.cell(r, 1).fill.solid(); t.cell(r, 1).fill.fore_color.rgb = RED_LT
    t.cell(r, 2).fill.solid(); t.cell(r, 2).fill.fore_color.rgb = GREEN_LT
icon_img(s, "x", 3.55, 4.25, 0.36, color="#C80019")
box(s, 3.95, 4.25, 2.0, 0.36, text="Static and blind", size=11, color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE)
icon_img(s, "check", 6.6, 4.25, 0.36, color="#0B7A42")
box(s, 7.0, 4.25, 2.4, 0.36, text="Short-lived and traced", size=11, color=GREEN_DK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
take(s, "Keep your IdP. Add the agent layer it was never built to cover.", icon="id-badge-2", y=4.8)

s = std("Three Vendor Categories Each Leave a Gap Open", "Competitive landscape", """
From the PRD competitive landscape.
IAM incumbents, for example Okta and Microsoft Entra, extend human and workload identity to agents. Strong identity. Limited runtime and prompt-layer enforcement.
Non-human identity (NHI) vendors focus on secrets and service-account hygiene. Good discovery. Weak on delegation chains and agent behavior.
AI security platforms, for example Zenity, cover agent posture, guardrails and prompt defenses. Strong AI context. Often SaaS-only and bound to a gateway.
The AccuKnox edge is on the next slide and in Part 6.
""")
mark("p2")
cats = [("id-badge-2", "IAM incumbents", "Okta, Microsoft Entra", "Strong identity", "Little runtime or prompt control"),
        ("key", "NHI vendors", "Secrets and service accounts", "Good discovery", "Weak on delegation and behavior"),
        ("robot", "AI security platforms", "Posture and guardrails", "Strong AI context", "Often SaaS-only, gateway-bound")]
for i, (ic, h, sub, plus, minus) in enumerate(cats):
    x = 0.40 + i * 3.1
    box(s, x, 1.22, 2.96, 3.25, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    icon_tile(s, ic, x + 0.2, 1.4, 0.6, bg=MUTE)
    box(s, x + 0.95, 1.4, 1.95, 0.34, text=h, size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    box(s, x + 0.95, 1.74, 1.95, 0.3, text=sub, size=10.5, color=MUTE, ml=0)
    box(s, x + 0.2, 2.35, 2.56, 0.85, fill=GREEN_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    icon_img(s, "check", x + 0.32, 2.58, 0.38, color="#0B7A42")
    box(s, x + 0.8, 2.35, 1.9, 0.85, text=plus, size=12, color=GREEN_DK, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    box(s, x + 0.2, 3.38, 2.56, 0.9, fill=RED_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    icon_img(s, "x", x + 0.32, 3.64, 0.38, color="#C80019")
    box(s, x + 0.8, 3.38, 1.9, 0.9, text=minus, size=12, color=RED, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
take(s, "No category ties identity, runtime enforcement and prompt defense together.", icon="puzzle", y=4.68)


# =================================================================== PART 3
divider("Part 3 · The AccuKnox solution", "Verify Every Agent and Every Action",
        "One control plane binds each agent to a verified identity, a scoped task and a full audit trail.",
        "shield-lock", "Part 3 shows the AccuKnox answer: the CISO mandate, the AgentZ control plane, the reference architecture and the request flow.")
mark("p3")

s = std("Zero Trust for Agents Is One Unbroken Chain", "The CISO mandate", """
From the PRD, the CISO mandate for Zero Trust agentic identity.
Verified human identity: the human principal who starts the chain, tied to the enterprise IdP.
Attested agent identity: cryptographic proof of what the workload is, not only a static secret it holds.
Scoped delegation: short-lived, audience-bound credentials with strict TTLs under 15 minutes.
Least-privilege actions: granular, per-action authorization on specific MCP tools and APIs.
The mandate: an unbroken chain from the first prompt to the final write.
""")
mark("p3")
chev = [("Verified" + BR + "human", "user", "The person who starts the chain, tied to your IdP", NAVY),
        ("Attested" + BR + "agent", "fingerprint", "Proof of what the workload is, not a secret it holds", PRIMARY),
        ("Scoped" + BR + "delegation", "key", "Short-lived credentials that expire in minutes", SECOND),
        ("Least" + BR + "privilege", "lock-access", "Per-action rights on named MCP tools and APIs", PURPLE)]
for i, (h, ic, b, c) in enumerate(chev):
    x = 0.40 + i * 2.3
    sp = box(s, x, 1.3, 2.42, 1.0, fill=c, shape=MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON,
             text=h, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ml=0.28, mr=0.06)
    icon_tile(s, ic, x + 0.75, 2.55, 0.62, bg=c)
    box(s, x + 0.1, 3.3, 2.05, 1.0, text=b, size=11.5, color=INK, align=PP_ALIGN.CENTER)
take(s, "The mandate: an unbroken chain from the first prompt to the final write.", icon="link", y=4.55)

s = std("AgentZ Puts Four Controls Around Every Agent", "The AccuKnox answer", """
From the PRD, the AccuKnox answer: the AgentZ control plane.
Identity and attestation: SPIFFE workload identity, with ModelKnox for model attestation.
Delegation and policy: OpenFGA relationship-based decisions and RFC 8693 delegation chains.
Execution guardrails: KubeArmor runtime enforcement and the stateful Prompt Firewall.
Credential broker: ephemeral access, scoped per task.
Positioning: evolutionary consolidation, not disruption. It sits as a control-plane layer directly above the existing infrastructure.
From accuknox.com/platform/agentz: agents never hold secrets. AgentZ injects scoped credentials at runtime and never stores them in the agent context, so a prompt injection cannot leak what the agent never receives.
""")
mark("p3")
dg(s, "control_plane", y=1.14, w=8.4)
take(s, "Agents never hold secrets. A prompt injection cannot leak what the agent never got.", icon="lock", y=4.55)

s = std("One Gateway Stands Between Every Agent and Every Tool", "Reference architecture", """
From the PRD reference architecture.
Humans and apps (users, services, schedulers) drive AI agents (agents, sub-agents, copilots).
Every agent call passes the AccuKnox Agent Identity Gateway: identity and attestation, credential broker, delegation, policy decision, prompt firewall, MCP gateway.
The gateway reaches tools and data: MCP servers, APIs, SaaS and databases.
Supporting services: the enterprise IdP (Okta, Entra) stays the system of record. Policy as code and the risk engine hold task scopes, risk tiers and simulation. The runtime sensor (KubeArmor eBPF) enforces at process, file and network level. Audit, lineage and AI-SOC write tamper-evident logs to SIEM and SOAR.
""")
mark("p3")
dg(s, "refarch", y=1.2, w=9.0)
take(s, "Your IdP stays the source of truth. AccuKnox adds the agent layer on top.", icon="id-badge-2", y=4.65)

s = std("Ten Capabilities Cover the Agent From Discovery to Kill Switch", "Capabilities at a glance", """
The top 10 capabilities from the PRD, in order. Part 4 covers each one.
01 Agent discovery and inventory: find and register every agent, sanctioned or shadow.
02 Verifiable agent identity: a unique, attested cryptographic identity per agent instance.
03 Secretless, short-lived credentials: remove static keys, broker short-lived tokens.
04 Delegation and on-behalf-of chains: carry the human principal through every hop.
05 Least-privilege, task-scoped authorization: just-in-time permissions that expire with the task.
06 Continuous runtime authorization: check every tool call, not only the login.
07 Tool and MCP governance: allowlist, pin and scope every tool and server.
08 Prompt, context and memory integrity: stop injected instructions from becoming actions.
09 Human-in-the-loop step-up: risk-tiered approvals for high-impact actions.
10 Audit, forensics and kill switch: full lineage, and revoke a rogue chain in seconds.
PRD priority: all P0 except 08 and 09, which are P1.
""")
mark("p3")
CAPS = [("radar", "Discovery and inventory"), ("fingerprint", "Verifiable identity"),
        ("key-off", "Secretless credentials"), ("route", "Delegation chains"),
        ("lock-access", "Task-scoped access"), ("activity", "Runtime authorization"),
        ("plug-connected", "Tool and MCP control"), ("message-exclamation", "Prompt and memory"),
        ("user-check", "Human approval"), ("history", "Audit and kill switch")]
for i, (ic, h) in enumerate(CAPS):
    x = 0.40 + (i % 5) * 1.86
    y = 1.22 + (i // 5) * 1.72
    box(s, x, y, 1.74, 1.58, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    box(s, x + 0.08, y + 0.06, 0.7, 0.4, text=f"{i+1:02d}", size=16, color=SECOND, bold=True)
    icon_tile(s, ic, x + 1.08, y + 0.14, 0.5, bg=PRIMARY if i < 5 else PURPLE)
    box(s, x + 0.1, y + 0.78, 1.56, 0.7, text=h, size=12, color=NAVY, bold=True)
take(s, "Phase 1 sees and stops agents. Later phases scope and govern them at scale.", icon="stairs-up", y=4.75)


# =================================================================== PART 4
divider("Part 4 · Ten capabilities in depth", "Each Capability Closes One Gap",
        "Two capabilities per slide: why it matters, what it does, and the AccuKnox building blocks behind it.",
        "list-details", "Part 4 walks the ten capabilities in pairs. User stories and acceptance criteria from the PRD sit in each slide's notes.")
mark("p4")


def cap_col(s, x, n, ic, name, why, items, blocks):
    box(s, x, 1.16, 4.5, 3.98, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    num_tile(s, x + 0.18, 1.3, n)
    box(s, x + 0.8, 1.3, 3.0, 0.5, text=name, size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    icon_img(s, ic, x + 3.9, 1.3, 0.48, color="#0046FF")
    box(s, x + 0.18, 1.92, 4.14, 0.62, text=why, size=11.5, color=NAVY, fill=LAV, italic=True,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, anchor=MSO_ANCHOR.MIDDLE, ml=0.14)
    for j, (bic, txt) in enumerate(items):
        y = 2.68 + j * 0.5
        icon_img(s, bic, x + 0.22, y + 0.04, 0.32, color="#16A55C")
        box(s, x + 0.66, y, 3.7, 0.44, text=txt, size=11.5, color=INK, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    box(s, x + 0.18, 4.22, 3.0, 0.26, text="ACCUKNOX BUILDING BLOCKS", size=9, color=GREY_TX, bold=True, ml=0)
    px = x + 0.18
    for b in blocks:
        w = 0.3 + 0.085 * len(b)
        pill(s, px, 4.54, w, 0.32, b, fill=NAVY, tcolor=WHITE, size=10)
        px += w + 0.1


def cap_slide(title, left, right, notes_text):
    s = std(title, "Capabilities " + left[0] + " and " + right[0], notes_text)
    mark("p4")
    cap_col(s, 0.4, *left)
    cap_col(s, 5.1, *right)
    return s


cap_slide("Know Every Agent and Prove Who It Is",
          ("01", "radar", "Discovery and inventory", "You cannot secure an agent you cannot see.",
           [("search", "Finds agents in Kubernetes, cloud AI, SaaS and laptops"),
            ("user-check", "Records owner, purpose, model, tools and data"),
            ("ghost", "Flags orphaned agents and unused rights")],
           ["AI-SPM", "CNAPP asset graph"]),
          ("02", "fingerprint", "Verifiable identity", "Shared accounts make every agent look the same.",
           [("id-badge-2", "One cryptographic ID per agent instance"),
            ("certificate", "Bound to image, model, config and runtime"),
            ("link", "Federates with Okta or Entra ID")],
           ["AgentZ", "KubeArmor", "ModelKnox"]),
          """
01 Agent discovery and inventory, P0.
Why it matters: agents are spun up by developers, embedded in SaaS copilots and run on laptops. Most never pass a security review.
User story: as a CISO, I want a live inventory of every agent, its owner and what it can reach, so shadow agents cannot operate unnoticed.
Core requirements: auto-discover agents across Kubernetes, cloud AI services, SaaS copilots and developer endpoints. Register each agent with owner, purpose, model, tools and data access. Flag orphaned agents (owner left, project ended) and unused privileges. Risk-score each agent from its access, exposure and behavior.
Acceptance criteria: every agent maps to a named human owner. Orphans alerted within 24 hours. Inventory exports to CMDB or SIEM.
02 Verifiable agent identity, P0.
Why it matters: Zero Trust starts with knowing exactly which agent, and which version, is asking.
User story: as an IAM architect, I want each agent instance to prove who it is and what code it runs before it gets anything.
Core requirements: issue a unique cryptographic identity per agent instance, SPIFFE and SVID-style workload identity. Bind identity to attestation: container image, model, prompt or config, and runtime hashes. Distinguish agent type, version and instance for policy and audit. Federate with the enterprise IdP rather than replacing it.
Acceptance criteria: no two agents share an identity. Every call carries a signed identity. A tampered image or config fails attestation.
""")

cap_slide("Give Agents No Standing Secrets and a Traceable Chain",
          ("03", "key-off", "Secretless credentials", "Static API keys are the top cause of over-privilege.",
           [("clock-hour-4", "Short-lived, audience-bound tokens only"),
            ("refresh", "Rotates and revokes through your vaults"),
            ("eye-off", "Keeps secrets out of prompts and memory")],
           ["AgentZ", "Secrets scanning"]),
          ("04", "route", "Delegation chains", "Agents act for people, and agents call agents.",
           [("user", "Carries the human through every hop"),
            ("arrow-down", "Scope can only shrink at each hop"),
            ("plug-connected", "Covers agent-to-agent and MCP calls")],
           ["AgentZ"]),
          """
03 Secretless, short-lived credentials, P0.
Why it matters: static API keys are the single most common way agents get over-privileged and compromised. They leak into prompts, logs and memory.
User story: as a platform engineer, I want my agents to get credentials on their own and only for as long as they need them.
Core requirements: broker short-lived, audience-bound tokens, with no long-lived keys in agent code or context. Automatic rotation and revocation with existing vaults and cloud IAM. Detect hard-coded or pasted secrets in prompts, code and agent memory. Keep secrets out of the model's context window entirely.
Acceptance criteria: default token TTL 15 minutes or less, configurable. Static-key use flagged on first sight. Zero secrets present in agent context.
04 Delegation and on-behalf-of chains, P0.
Why it matters: without an unbroken chain you get confused-deputy attacks and no accountability.
User story: as a SOC analyst, I want every agent action tied to the human who started it, through every hop.
Core requirements: carry user to agent to sub-agent delegation with OAuth 2.0 Token Exchange (RFC 8693) actor claims. Attenuate scope at every hop, so a delegate never exceeds its delegator. Record consent and delegation grants with expiry. Support agent-to-agent (A2A) and MCP call paths.
Acceptance criteria: 100 percent of actions attributable to a human principal. Downstream scopes are a subset of upstream scopes, enforced. Delegation revocable mid-chain.
""")

cap_slide("Authorize Each Action, Not Only the Login",
          ("05", "lock-access", "Task-scoped access", "Agents get standing access just in case.",
           [("clock-hour-4", "Just-in-time grants that expire with the task"),
            ("code", "Policy as code with data, time and risk"),
            ("adjustments", "Right-sizing from usage, dry run first")],
           ["AgentZ", "AI-SPM"]),
          ("06", "activity", "Runtime authorization", "A trusted agent can be hijacked one prompt later.",
           [("scale", "A decision on every tool and API call"),
            ("chart-dots", "Per-agent baseline blocks odd behavior"),
            ("cpu", "Enforced at process, file and network level")],
           ["KubeArmor", "AgentZ"]),
          """
05 Least-privilege, task-scoped authorization, P0.
Why it matters: least privilege has to be per task, just-in-time and self-expiring.
User story: as an IAM architect, I want permissions granted for a task and gone when it ends, without writing policy by hand for every agent.
Core requirements: just-in-time, task-scoped grants that expire on completion or timeout. Policy as code (for example OPA or Cedar) with attribute and relationship context: data sensitivity, time, risk. Right-sizing recommendations from observed usage. Simulation or dry-run mode before enforcement.
Acceptance criteria: grants auto-expire at task end. Unused privileges surfaced weekly. Policies testable in simulation.
06 Continuous runtime authorization, P0.
Why it matters: authorization must happen on every action, all the time.
User story: as a CISO, I want every tool and API call checked in real time, with enforcement at the runtime and not only at the gateway.
Core requirements: real-time policy decision on every tool call, API request and data access. Behavioral baselining per agent, with anomalies that trigger step-down or block. Runtime enforcement at process, file and network level (eBPF and LSM). Block, alert and audit-only modes per policy.
Acceptance criteria: p99 decision latency target to be confirmed, under 25 ms proposed. Out-of-baseline actions blocked in enforce mode. Works in air-gapped clusters.
""")

cap_slide("Control the Tools Agents Use and the Text They Read",
          ("07", "plug-connected", "Tool and MCP control", "A poisoned tool turns a helper into an attacker.",
           [("list-check", "Registry and allowlist for MCP servers"),
            ("certificate", "Signed, pinned manifests catch rug pulls"),
            ("adjustments", "Per-tool scopes and argument limits")],
           ["AgentZ", "AI-SPM"]),
          ("08", "message-exclamation", "Prompt and memory", "Injected text hijacks agents without stealing keys.",
           [("history", "Stateful firewall tracks intent across turns"),
            ("file-search", "Inspects retrieved content and tool output"),
            ("eye-off", "Catches memory poisoning, redacts data")],
           ["Prompt Firewall", "DSPM"]),
          """
07 Tool and MCP governance, P0.
Why it matters: tools are how agents touch the world. A poisoned tool description or rogue MCP server turns a helpful agent into an attacker.
User story: as a platform engineer, I want agents to use only approved tools, and to know when a tool changes underneath them.
Core requirements: registry and allowlist for MCP servers and tools, block unapproved ones. Sign and pin tool manifests, detect description changes (rug pulls). Per-tool scopes and argument-level constraints. An MCP gateway that enforces authorization on every call.
Acceptance criteria: unapproved MCP servers blocked. Tool definition changes require re-approval. Per-tool usage visible per agent.
08 Prompt, context and memory integrity, P1.
Why it matters: injected instructions in documents, web pages and tool output are how attackers hijack agents without touching credentials.
User story: as a security architect, I want injected instructions caught before they turn into tool calls or data exfiltration.
Core requirements: stateful prompt firewall that tracks intent across multi-turn sessions. Inspect retrieved content and tool outputs for indirect injection. Detect memory poisoning in long-term agent memory. Classify and redact sensitive data in prompts and responses.
Acceptance criteria: injection blocked before tool execution. Sensitive data redacted per policy. Session context retained across turns.
""")

cap_slide("Keep People in the Loop and Stop a Rogue Chain Fast",
          ("09", "user-check", "Human approval", "Some actions are too big for autonomy alone.",
           [("stairs-up", "Risk tiers by action, data and blast radius"),
            ("message", "Approve in Slack, Teams, email or mobile"),
            ("ban", "Unanswered requests deny by default")],
           ["AgentZ", "AI-SOC"]),
          ("10", "history", "Audit and kill switch", "When it goes wrong, you need the story and a stop.",
           [("route", "Tamper-evident lineage to the data touched"),
            ("radar", "Exports to SIEM and SOAR, AI-SOC views"),
            ("player-stop", "Revokes an agent, a version or a chain")],
           ["AI-SOC", "AgentZ", "KubeArmor"]),
          """
09 Human-in-the-loop step-up, P1.
Why it matters: some actions are too consequential for autonomy alone, such as moving money, deleting data, changing privileges or contacting customers.
User story: as a CISO, I want high-risk agent actions to pause for a human, without slowing down routine work.
Core requirements: risk tiers per action type, data class and blast radius. Asynchronous approval through Slack, Teams, email or mobile (CIBA-style). Timeouts default to deny, with break-glass and full audit. Tune thresholds from approval history to cut noise.
Acceptance criteria: tiers configurable per agent and action. Unanswered approvals deny by default. Approval decisions logged with the approver.
10 Audit, forensics and kill switch, P0.
Why it matters: when something goes wrong, you need to rebuild exactly what happened and stop it at once across every agent in the chain.
User story: as a SOC analyst, I want a full lineage of an incident and a single control to quarantine a rogue agent chain.
Core requirements: tamper-evident lineage from request to delegation to tool call to data touched. Export to SIEM and SOAR, with native AI-SOC investigation views. Kill switch: revoke or quarantine an agent, a version or an entire chain. Compliance mapping to the EU AI Act, NIST AI RMF, ISO 42001 and SOC 2.
Acceptance criteria: full chain rebuildable per action. Revocation propagates in a time to be confirmed, under 60 seconds proposed. Evidence packs export for auditors.
""")

s = std("Every Request Passes Five Checkpoints", "How one request flows", """
From the PRD, the complete agentic Zero Trust blueprint: eager auth, token exchange and CEL RBAC.
1. Connect: the enterprise client registers and authenticates once at connect time (eager auth).
2. Invoke: the autonomous agent tries to call a tool.
3. Exchange: the security token service handles impersonation, delegation and on-behalf-of exchange.
4. Authorize: a CEL RBAC policy filters the call. Example rule: request.auth.claims.act.sub contains 'finance-agent'.
5. Execute: a scoped, audited token reaches the SaaS MCP server.
Steps 2 to 5 run inside the trust domain and security perimeter.
""")
mark("p4")
dg(s, "blueprint", y=1.18, w=9.0)
take(s, "The agent never carries the user's token to the tool. It gets a scoped one per call.", icon="arrows-exchange", y=4.4)

s = std("On-Behalf-Of Exchange Keeps the User on Every Hop", "Example with Entra ID", """
From the PRD: set up OAuth 2.0 on-behalf-of token exchange with Entra ID.
1. User and client: the developer signs in from an MCP client with authorization code and PKCE.
2. Entra ID, the IdP and authorization server: OIDC authentication with OAuth 2.1 and PKCE, custom auth server, signing keys and JWKS, scopes and claims, access policies, trusted servers. It issues the broker JWT.
3. AccuKnox Agent Fabric, the agent gateway. Inbound, it validates the JWT: signature, issuer, audience, expiration, scope. Outbound, it runs the on-behalf-of token exchange (RFC 8693) and injects the delegated JWT.
4. The protected MCP resource gets only the delegated JWT.
Planner agent broker JWT: audience api://planning-agent-broker, scope Planner Agent Broker Invoke, subject the user.
MCP server delegated JWT: audience api://MCPServer, scope MCP Server Invoke, subject the original user, client Broker OBO Client.
Rule: never forward the planner agent broker JWT downstream.
""")
mark("p4")
dg(s, "obo", y=1.14, w=8.5)
take(s, "The tool sees the original user and a narrow scope, never the broker token.", icon="shield-check", y=4.68)

s = std("Authorization Follows the Whole Caller Chain", "Worked example", """
From the technical deck. When a request reaches the AI gateway, it may have passed several agents. The decision depends on which agents were in the path. This is unique to agentic AI.
Do not assume the principal's permission for the AI agent.
Decision inputs: the principal (the user the request acts for), the upstream caller sequence (every agent the request passed), and the target resource (the MCP server or tool).
Example rules from the multi-cloud scenario. The user has read and write rights on the GitHub repository.
Target GitHub MCP server, principal User-1, upstream Coder agent: read on all repos.
Target GitHub MCP server, principal User-1, upstream Remediation agent then Coder agent: read-only repos, read-write PRs.
Target ServiceNow MCP server, principal User-1, upstream Remediation agent then Security agent: read-write tickets.
Lineage is kept as a signed delegation-and-receipt system, not a mutable HTTP header stack. Each hop binds to an authenticated workload identity.
""")
mark("p4")
dg(s, "chain", y=1.18, w=9.0)
take(s, "Same user, different agent path, different rights. OpenFGA makes that call.", icon="route", y=4.55)

s = std("Four Layers Carry One Identity Proof From the Workload Up", "The agentic mesh stack", """
From the PRD and the technical deck, the agentic mesh stack. Read it bottom up.
Layer 1, workload identity anchor: the SPIRE server issues the SVID that every layer above trusts.
Layer 2, ambient service mesh: every hop authenticated without a sidecar per pod, with L4 mTLS.
Layer 3, translation and control: the MCP gateway and security token service route MCP auth and mint single-hop JWTs per call.
Layer 4, enterprise identity: dual IdP setup with Keycloak for internal OIDC and Microsoft Entra ID for the cloud domain.
An agent proves what it is at Layer 1, and every layer above reuses that proof.
""")
mark("p4")
dg(s, "mesh", y=1.16, w=8.6)
take(s, "An agent proves what it is once, at Layer 1. Every layer above reuses that proof.", icon="stack-2", y=4.6)

s = std("Deploy Anywhere and Keep the Identity Stack You Have", "Platform", """
From the PRD non-functional requirements.
Deployment: SaaS, customer VPC, private cloud and fully air-gapped, with the same feature set and no phone-home.
Standards: SPIFFE and SVID, OAuth 2.1, RFC 8693 token exchange, OpenID, MCP authorization spec.
Integrations: Okta, Entra, Ping. HashiCorp Vault and cloud KMS. Splunk and Microsoft Sentinel. Slack and Teams.
Privacy and sovereignty: prompts and data stay in the customer's boundary, with configurable retention and redaction.
Scale target: tens of thousands of agent instances and millions of tool calls per day per tenant. Performance target: p99 policy latency under 25 ms proposed, no single point of failure. Both are PRD targets, not shipped numbers.
""")
mark("p4")
deps = [("cloud", "SaaS"), ("cloud-lock", "Customer VPC"), ("server-2", "Private cloud"), ("wifi-off", "Air-gapped")]
for i, (ic, h) in enumerate(deps):
    x = 0.40 + i * 2.33
    box(s, x, 1.2, 2.18, 1.1, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    icon_tile(s, ic, x + 0.16, 1.42, 0.62, bg=PRIMARY if i < 3 else NAVY)
    box(s, x + 0.9, 1.42, 1.22, 0.62, text=h, size=13, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
label(s, CX, 2.45, 4, "OPEN STANDARDS", size=10)
px = CX
for st in ["SPIFFE and SVID", "OAuth 2.1", "RFC 8693", "OpenID Connect", "MCP authorization"]:
    w = 0.36 + 0.088 * len(st)
    pill(s, px, 2.76, w, 0.36, st, fill=LAV, tcolor=NAVY, size=11)
    px += w + 0.14
label(s, CX, 3.3, 4, "WORKS WITH YOUR STACK", size=10)
logo_row(s, [("si-okta", "Okta"), ("si-microsoftazure", "Entra ID"), ("vault-original", "Vault"),
             ("si-splunk", "Splunk"), ("si-microsoft", "Sentinel"), ("slack-original", "Slack"),
             ("si-microsoftteams", "Teams")], CX, 3.62, CW, size=0.48, label_size=10.5)
take(s, "Same features in every deployment, no phone-home. Prompts and data stay with you.", icon="lock", y=4.78)

s = std("Identity Plugs Into the Full AccuKnox AI Security Platform", "Beyond identity", """
Agent identity is one layer of the AccuKnox AI security platform. The other layers share the same control plane and inventory.
AI-SPM and AIBOM: an AI bill of materials across managed and unmanaged AI assets, native integration with most model providers, CycloneDX v1.7 format. Source: technical deck, AI asset inventory slide.
Prompt Firewall: covers SaaS AI apps (Claude.ai, chatgpt.com, gemini.google.com, Microsoft Copilot, GitHub Copilot), cloud agents (Azure Copilot Studio and Power Apps, AWS Bedrock AgentCore, GCP Agents Platform), dev agents (Claude Code, Codex, Antigravity) and local agents (LangGraph, n8n and others). Five integration modes: browser plugin, gateway proxy, SDK or proxy mode, cloud SDK or API gateway, and Purview for Microsoft Copilot agents.
KubeArmor sandbox: process, filesystem, credential and network isolation for agents and MCP servers, at the kernel with eBPF.
AI-SOC: triage and investigation across cloud, identity and AI workloads.
DSPM: classifies the data agents can reach.
CIEM: cloud IAM users, roles and groups with risk class. Covered next.
""")
mark("p4")
plat = [("radar", "AI-SPM and AIBOM", "Every model, agent and MCP server"),
        ("message-exclamation", "Prompt Firewall", "Stops injection and data leaks"),
        ("box", "KubeArmor sandbox", "Kernel-level limits per agent"),
        ("activity", "AI-SOC", "Investigates agent incidents"),
        ("database", "DSPM", "Knows the data agents touch"),
        ("users", "CIEM", "Cloud identities and their risk")]
for i, (ic, h, b) in enumerate(plat):
    x = 0.40 + (i % 3) * 3.1
    y = 1.18 + (i // 3) * 1.32
    icon_card(s, x, y, 2.96, 1.2, ic, h, b, hsize=13, bsize=11, accent=PRIMARY if i != 5 else PURPLE)
label(s, CX, 3.86, 6, "PROMPT FIREWALL COVERS SAAS, CLOUD, DEV AND LOCAL AGENTS", size=10)
logo_row(s, [("si-openai", "ChatGPT"), ("si-claude", "Claude"), ("si-googlegemini", "Gemini"),
             ("si-githubcopilot", "Copilot"), ("amazonwebservices-original-wordmark", "AgentCore"),
             ("si-langchain", "LangGraph"), ("si-n8n", "n8n")], CX, 4.16, CW, size=0.44, label_size=10.5)


# =================================================================== PART 5, CIEM
divider("Part 5 · CIEM in the product", "See Every Cloud Identity and the Access It Holds",
        "AccuKnox CIEM lists human and machine identities across four clouds and traces each one to its policies.",
        "fingerprint", "Part 5 shows AccuKnox CIEM screens from help.accuknox.com. CIEM covers cloud IAM identities today, the same machine identities agents run on.")
mark("p5")

s = std("CIEM Lists Every Cloud Identity With Its Risk", "CIEM · Identity list", """
Source: help.accuknox.com, Cloud Infrastructure Entitlement Management (docs/getting-started/ciem-overview.md).
CIEM lists the users, groups and roles in AWS, GCP, Azure and Oracle Cloud Infrastructure accounts. For each identity it shows the attached policies, the risk classes those policies carry, and the last time the identity was used.
Each identity is classed human or machine. A CI/CD pipeline account is a machine identity.
Risk class examples: privilege escalation, data exfiltration, credentials exposure, infrastructure modification.
Lifecycle columns: age in days, created at, created by, last used, days since last used.
Where to find it: Identities, then CIEM, in the left navigation.
Status: the help page marks CIEM as coming soon in a newer AccuKnox version. Screens come from the current build and can change. CIEM supports standalone cloud accounts today, not organization accounts.
Screenshot: ciem-identity-list.png, account values redacted in the source.
""")
mark("p5")
image_fit(s, shot("ciem-identity-list.png", crop=(110, 110, 2560, 1254)), CX, 1.2, 6.15, 3.4, caption="Identities > CIEM, list view filtered to AWS", capsize=9.5)
x0 = 6.75
logo_row(s, [("amazonwebservices-original-wordmark", "AWS"), ("googlecloud-original", "GCP"),
             ("azure-original", "Azure"), ("oracle-original", "OCI")], x0, 1.2, 2.85, size=0.42, label_size=10)
facts = [("users", "Users, groups and roles in one list"),
         ("robot", "Each one marked human or machine"),
         ("alert-triangle", "Risk class from attached policies"),
         ("clock-hour-4", "Last used and age show stale access")]
for j, (ic, t) in enumerate(facts):
    y = 2.05 + j * 0.62
    icon_tile(s, ic, x0, y, 0.44, bg=PRIMARY if j != 2 else RED)
    box(s, x0 + 0.56, y - 0.04, 2.3, 0.52, text=t, size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE, ml=0)
pill(s, x0, 4.6, 1.5, 0.3, "Next release", fill=LAV, tcolor=NAVY, size=10)
take(s, "Privilege escalation and data exfiltration risk shows on each identity row.", icon="alert-triangle", y=4.98, h=0.4)

s = std("The Access Graph Shows Where Each Permission Comes From", "CIEM · Identity detail", """
Source: help.accuknox.com CIEM overview.
Select an identity in the list. A panel opens with four tabs: Overview, Access Graph, Related Identities and Raw Information.
Access Graph: the identity sits at the top, its groups in the middle, and the policies each group grants at the bottom. A policy attached straight to the identity sits in the same row as the groups. It shows which permissions reach the identity and through which group.
Overview: identity details, total findings by severity (Critical, High, Medium, Low), and the attached policies. Expand a policy row to read the policy document as JSON. Policy types: cloud managed, inline policy, user managed.
Screenshots: ciem-identity-graph.png and ciem-identity-overview.png, cropped to the side panel. Identity names are redacted in the source.
""")
mark("p5")
image_fit(s, shot("ciem-identity-graph.png", crop=(790, 300, 2560, 1254)), CX, 1.2, 6.9, 3.45, caption="Access Graph tab: one user, its groups, and the policies they grant", capsize=9.5)
x0 = 7.5
for j, (ic, t) in enumerate([("hierarchy-3", "Identity to group to policy"), ("file-search", "Read each policy as JSON"), ("alert-triangle", "Findings by severity"), ("users", "Related identities in the account")]):
    y = 1.25 + j * 0.85
    icon_tile(s, ic, x0, y, 0.44, bg=PURPLE)
    box(s, x0 + 0.54, y - 0.06, 1.6, 0.6, text=t, size=11, color=INK, anchor=MSO_ANCHOR.MIDDLE, ml=0)
take(s, "Trace an excess permission to the exact group or policy that granted it.", icon="route", y=4.92, h=0.42)

s = std("One Graph Maps Every User, Group, Role and Policy", "CIEM · Organization graph", """
Source: help.accuknox.com CIEM overview.
Select the graph icon at the top right of Identities, CIEM. Filter by cloud provider and cloud account name.
The cloud account node sits at the top, with one node each for users, groups, roles and policies. The badge on each node is the count. In the sample AWS account on the help page, the badges read 77 users, 22 groups, 637 roles and 504 policies. Select a node to expand its members, then a member to open its access graph.
Bridge to the agentic story: roles far outnumber users in the sample account, the same machine-identity pattern behind the 82 to 1 ratio. Agentic identity security extends the same visibility and control to AI agents.
Screenshot: ciem-org-graph.png.
""")
mark("p5")
image_fit(s, shot("ciem-org-graph.png", crop=(100, 280, 2560, 1254)), CX, 1.2, 6.15, 3.3, caption="Organization graph for one AWS account", capsize=9.5)
x0 = 6.8
for j, (n, l) in enumerate([("77", "users"), ("22", "groups"), ("637", "roles"), ("504", "policies")]):
    x = x0 + (j % 2) * 1.42
    y = 1.2 + (j // 2) * 1.25
    stat(s, x, y, 1.32, n, l, color=PURPLE if n == "637" else PRIMARY, nsize=24, lsize=11, h=1.1)
box(s, x0, 3.75, 2.8, 0.7, text="Sample AWS account. Roles outnumber users 8 to 1.", size=11, color=MUTE)
take(s, "CIEM covers cloud identities today. Agentic identity extends it to every AI agent.", icon="robot", y=4.92, h=0.42)


# =================================================================== PART 6
divider("Part 6 · Why AccuKnox", "Identity and Runtime Enforcement in One Platform",
        "Four differentiators, a category comparison, the outcomes you can prove, and the live demo.",
        "rocket", "Part 6 closes on the key differentiators and the proof.")
mark("p6")

s = std("Four Differentiators Set AccuKnox Apart", "Key differentiators", """
From the PRD, the AccuKnox edge.
1. Identity plus runtime enforcement in one platform, through KubeArmor eBPF. Decisions are enforced at process, file and network level, not only at a gateway.
2. Stateful prompt firewall, patent pending. It tracks intent across multi-turn sessions.
3. Air-gapped and sovereign deployment. SaaS, customer VPC, private cloud or fully air-gapped, with the same feature set and no phone-home.
4. CNAPP, AI-SPM and DSPM context. Agent identity decisions use the cloud, AI asset and data context AccuKnox already collects.
Built on open standards: SPIFFE for identity and OpenFGA for fine-grained authorization (accuknox.com/solutions/agentic-ai-security).
""")
mark("p6")
diffs = [("cpu", "Identity plus runtime", "KubeArmor eBPF enforces each decision at process, file and network level"),
         ("message-exclamation", "Stateful prompt firewall", "Patent pending. Tracks intent across multi-turn sessions"),
         ("wifi-off", "Air-gapped and sovereign", "Same features in SaaS, VPC, private cloud or offline"),
         ("stack-2", "Cloud, AI and data context", "CNAPP, AI-SPM and DSPM feed every identity decision")]
for i, (ic, h, b) in enumerate(diffs):
    x = 0.40 + (i % 2) * 4.66
    y = 1.2 + (i // 2) * 1.6
    box(s, x, y, 4.54, 1.45, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    box(s, x, y, 0.1, 1.45, fill=PRIMARY)
    icon_tile(s, ic, x + 0.28, y + 0.22, 0.66)
    box(s, x + 1.12, y + 0.16, 3.3, 0.42, text=h, size=15, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
    box(s, x + 1.12, y + 0.6, 3.3, 0.8, text=b, size=11.5, color=MUTE, ml=0)
take(s, "Built on open standards: SPIFFE for identity and OpenFGA for authorization.", icon="certificate", y=4.48)

s = std("AccuKnox Covers What Each Category Leaves Out", "Category comparison", """
Category-level view built from the PRD competitive landscape. It compares vendor categories, not named products.
IAM incumbents: strong identity, limited runtime and prompt-layer enforcement.
NHI vendors: good discovery, weak on delegation chains and agent behavior.
AI security platforms: strong AI context, often SaaS-only and gateway-bound.
Partial means some vendors in the category offer it, or it works only at one layer such as a gateway.
Validate each cell with product marketing before this slide goes to a public page.
""")
mark("p6")
Y, P, N = "Yes", "Partial", "No"
rows = [["Capability", "IAM incumbents", "NHI vendors", "AI security", "AccuKnox"],
        ["Verified agent identity", Y, P, P, Y],
        ["Agent and NHI discovery", P, Y, Y, Y],
        ["Delegation chain control", P, N, P, Y],
        ["Kernel runtime enforcement", N, N, N, Y],
        ["Prompt-layer defense", N, N, Y, Y],
        ["Air-gapped deployment", P, P, N, Y]]
t = table(s, CX, 1.2, CW, 3.4, rows, colw=[2.8, 1.6, 1.6, 1.6, 1.6], hsize=12, bsize=11.5,
          aligns=[PP_ALIGN.LEFT] + [PP_ALIGN.CENTER] * 4)
tint_cells(t, {Y: (GREEN_LT, GREEN_DK), P: (LAV, PURPLE), N: (RED_LT, RED)})
t.cell(0, 4).fill.solid(); t.cell(0, 4).fill.fore_color.rgb = PRIMARY
take(s, "Only AccuKnox ties identity, kernel enforcement and prompt defense together.", icon="trophy", y=4.78)

s = std("Five Outcomes a CISO Can Prove to the Board", "Outcomes", """
Outcomes drawn from the PRD goals and acceptance criteria.
Every agent known and owned: every agent maps to a named human owner, and orphans are flagged.
No static secrets on agents: short-lived, task-scoped credentials by default.
Every action tied to a human: delegation chains carry the human principal through every hop.
A rogue chain stopped from one control: revoke or quarantine an agent, a version or an entire chain.
Evidence for auditors: control mapping to the EU AI Act, NIST AI RMF, ISO 42001 and SOC 2, with evidence packs.
The PRD proposes numeric targets for these outcomes. They sit on a hidden backup slide and are not commitments.
""")
mark("p6")
outs = [("eye", "Every agent known and owned"), ("key-off", "No static secrets on agents"),
        ("user-check", "Every action tied to a human"), ("player-stop", "A rogue chain stopped at once"),
        ("gavel", "Evidence for EU AI Act, ISO 42001")]
for i, (ic, h) in enumerate(outs):
    x = 0.40 + i * 1.86
    box(s, x, 1.25, 1.74, 2.6, fill=GREEN_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    icon_tile(s, ic, x + 0.5, 1.5, 0.74, bg=GREEN_DK)
    icon_img(s, "circle-check", x + 1.3, 1.36, 0.32, color="#16A55C")
    box(s, x + 0.1, 2.45, 1.54, 1.2, text=h, size=12.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
box(s, CX, 4.1, CW, 0.9, text="Trust nothing. Verify every agent.", size=24, color=WHITE, bold=True,
    fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

s = std("See Multi-Cloud Agent Chains Run Under Policy, Live", "Demo and next steps", """
Demo scope, from the technical deck.
Showcase authorization rules enforcement.
Showcase cascaded multi-cloud agent calls and the authorization rules based on them. The Remediation and Coder agents run on AWS AgentCore, and the Security agent runs on Google ADK. Every call passes the on-prem AI gateway in Kubernetes.
Showcase SPIFFE for identity.
Showcase visibility across AWS, GCP and on-prem agents from the AccuKnox console.
Next steps for the customer: pick the first agent workflows to protect, run a design-partner pilot on them, and map the controls to your frameworks.
Watch: AgentZ launch film youtube.com/watch?v=mAzLWcr59g0 and the Kubernetes live demo with Rahul Jadhav, CTO, youtube.com/watch?v=dX1sO3sOoFI.
""")
mark("p6")
label(s, CX, 1.18, 4, "WHAT THE DEMO SHOWS", size=10)
demo = [("scale", "Authorization rules enforced per call"),
        ("route", "Agent chains across AWS and Google Cloud"),
        ("fingerprint", "SPIFFE identity for every agent"),
        ("eye", "One console view across clouds")]
for j, (ic, t) in enumerate(demo):
    y = 1.5 + j * 0.62
    icon_tile(s, ic, CX, y, 0.46)
    box(s, CX + 0.6, y, 3.9, 0.46, text=t, size=12, color=INK, anchor=MSO_ANCHOR.MIDDLE, ml=0)
logo_row(s, [("amazonwebservices-original-wordmark", "AgentCore"), ("googlecloud-original", "Google ADK"),
             ("kubernetes-original", "On-prem K8s")], CX, 4.05, 4.4, size=0.42, label_size=10)
box(s, 5.1, 1.18, 4.5, 3.75, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
box(s, 5.35, 1.32, 4.0, 0.4, text="Next steps", size=16, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
for j, t in enumerate(["Pick the first agent workflows to protect", "Run a design-partner pilot on them", "Map the controls to your frameworks"]):
    y = 1.95 + j * 0.85
    box(s, 5.35, y, 0.5, 0.5, text=str(j + 1), size=15, color=WHITE, bold=True, fill=PRIMARY, shape=MSO_SHAPE.OVAL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, ml=0, mr=0)
    box(s, 6.0, y - 0.06, 3.4, 0.62, text=t, size=12.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, ml=0)
icon_img(s, "player-play", 5.35, 4.38, 0.4, color="#FFFFFF")
box(s, 5.85, 4.36, 3.6, 0.44, text="Book a demo at accuknox.com/demo", size=11.5, color=NAVY_TXT, anchor=MSO_ANCHOR.MIDDLE, ml=0)


# =================================================================== closing
closing_slide(prs)


# =================================================================== hidden backup
def backup(title, kicker, cols, notes_text):
    s = std(title, kicker, notes_text)
    hide_slide(s)
    w = CW / len(cols) - 0.15
    for i, (h, ic, color, items) in enumerate(cols):
        x = CX + i * (w + 0.15)
        box(s, x, 1.2, w, 3.95, fill=WHITE, line=GREY_BD, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
        icon_tile(s, ic, x + 0.16, 1.32, 0.46, bg=color)
        box(s, x + 0.72, 1.32, w - 0.8, 0.46, text=h, size=14, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE, ml=0)
        bullets(s, x + 0.18, 1.98, w - 0.3, 3.1, items, size=13, mcolor=color, gap=8)
    return s


backup("Backup: The PRD Sets Five Goals and Four Non-Goals", "Backup · Scope", [
    ("Goals", "target", GREEN_DK, ["Every agent has a verifiable, owned identity", "No static secrets by default",
                                   "Every action authorized and tied to a human", "One policy model for people, workloads, agents",
                                   "Deploy in SaaS, private cloud or air-gapped"]),
    ("Non-goals for v1", "ban", RED, ["Replacing Okta or Entra, AccuKnox federates", "Building an agent framework",
                                      "Judging model quality or accuracy", "Consumer agent identity"])],
    """Hidden backup. From the PRD scope slide, goals and non-goals for v1. Internal planning material, keep it out of customer sessions unless asked.""")

backup("Backup: Proposed Targets for Design Partners to Validate", "Backup · Success metrics", [
    ("Coverage", "eye", PRIMARY, ["95% or more of running agents discovered and owned", "100% of P0 agents free of static secrets",
                                  "100% of actions tied to a human"]),
    ("Speed", "clock-hour-4", PURPLE, ["Under 60 s mean time to revoke a rogue chain", "Under 25 ms p99 inline policy decision",
                                       "Under 5% of human approvals judged unnecessary"])],
    """Hidden backup. From the PRD success metrics. These targets are proposals for validation with design partners, not commitments. Never present them as shipped performance.""")

backup("Backup: Three Phases Deliver the Ten Capabilities", "Backup · Roadmap", [
    ("Phase 1, 0 to 90 days", "eye", PRIMARY, ["See it and stop it", "01 Discovery and inventory", "02 Verifiable identity",
                                               "03 Secretless credentials", "07 Tool and MCP control", "10 Audit and kill switch"]),
    ("Phase 2, 90 to 180 days", "lock-access", PURPLE, ["Scope it, verify every action", "04 Delegation chains",
                                                        "05 Task-scoped access", "06 Runtime authorization", "08 Prompt and memory"]),
    ("Phase 3, 180 to 360 days", "stairs-up", GREEN_DK, ["Govern it at scale", "09 Human approval",
                                                         "Advanced approval tuning", "Cross-tenant agent federation", "Compliance evidence packs"])],
    """Hidden backup. From the PRD phased roadmap. Internal plan, dates count from project start and are not customer commitments.""")

backup("Backup: Open Risks and Questions in the PRD", "Backup · To resolve", [
    ("Risks", "alert-triangle", RED, ["Inline latency slows agents and drives bypass", "MCP auth, A2A and agent ID specs still move",
                                      "Approval fatigue erodes human review", "IdP vendors bundle good-enough agent identity"]),
    ("Open questions", "help", PURPLE, ["Which 3 to 5 design partners and frameworks first", "Gateway first, sidecar first, or both at GA",
                                        "Pricing unit: per agent, action or workload", "How deep to go with Okta and Entra agent ID",
                                        "Minimum viable air-gapped package"])],
    """Hidden backup. From the PRD risks and open questions. Internal only.""")


# =================================================================== TOC ranges
def rng(k):
    a, b = min(SECTIONS[k]), max(SECTIONS[k])
    return f"Slides {a} to {b}"


for tb, k in zip(TOC_RANGES, ["p1", "p2", "p3", "p4", "p5", "p6"]):
    r = tb.text_frame.paragraphs[0].runs[0]
    r.text = rng(k)

prs.save(OUT)
print(OUT, len(prs.slides), "slides")
