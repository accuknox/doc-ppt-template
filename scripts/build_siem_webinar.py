# -*- coding: utf-8 -*-
"""
AccuKnox SIEM webinar deck, 12 slides, 20 minute moderated session.
Moderator: Atharva Shah. Speaker: Aditya Raj.
Content sourced from AccuKnox_SIEM_Proposal.pdf and the July 2026 SIEM deck.

Render:  powershell -File scripts/render.ps1 -Pptx <OUT> -Out <dir>
"""
import os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
SRC  = os.path.join(HERE, "..", "PPT Template.pptx")
OUT  = os.path.join(HERE, "..", "output", "AccuKnox_SIEM_Webinar.pptx")

NAVY     = RGBColor(0x11, 0x20, 0x6D)
PRIMARY  = RGBColor(0x00, 0x46, 0xFF)
SECOND   = RGBColor(0x64, 0x64, 0xFF)
PURPLE   = RGBColor(0x4D, 0x4D, 0xD9)
RED      = RGBColor(0xC8, 0x00, 0x19)
GREEN    = RGBColor(0x16, 0xA5, 0x5C)
GREEN_DK = RGBColor(0x0B, 0x7A, 0x42)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
INK      = RGBColor(0x1B, 0x22, 0x3B)
MUTE     = RGBColor(0x5A, 0x63, 0x7D)
GREY_BD  = RGBColor(0xC4, 0xCC, 0xDE)
NAVY_TXT = RGBColor(0xB8, 0xC4, 0xE8)
FONT     = "Space Grotesk"
ICF      = "Segoe MDL2 Assets"

IC = {'lock':0xE72E,'shield':0xE83D,'cloud':0xE753,'globe':0xE774,'pulse':0xE9D9,
      'doc':0xE8A5,'person':0xE7EE,'gear':0xE713,'gears':0xE9F5,'key':0xE8D7,
      'bolt':0xE945,'warn':0xE7BA,'check':0xE930,'flag':0xE7C1,'bulb':0xEB50,
      'server':0xE968,'devices':0xE977,'chip':0xE964,'search':0xE721,'clock':0xE81C,
      'star':0xE734,'books':0xE8F1,'cert':0xEB95,'eye':0xE890,'list':0xE71D,
      'rocket':0xE7A7,'wrench':0xE90F,'mail':0xE715,'chat':0xE8BD}
def G(n): return chr(IC[n])

CX, CW = 0.40, 9.20

def _set_title(slide, text, color=None):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            ph.text_frame.paragraphs[0].text = text
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.name = FONT
                if color is not None: r.font.color.rgb = color
            return ph

def _blank_footer(slide):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in (10, 11):
            ph.text_frame.paragraphs[0].text = ""

def box(slide, x, y, w, h, text="", size=12, color=INK, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, fill=None, line=None,
        line_w=0.75, shape=MSO_SHAPE.RECTANGLE, radius=None, italic=False,
        wrap=True, ml=0.08, mr=0.08, mt=0.04, mb=0.04, font=FONT):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if radius is not None:
        try: sp.adjustments[0] = radius
        except Exception: pass
    tf = sp.text_frame; tf.word_wrap = wrap
    tf.margin_left=Inches(ml); tf.margin_right=Inches(mr)
    tf.margin_top=Inches(mt); tf.margin_bottom=Inches(mb)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    if text:
        r = p.add_run(); r.text = text
        r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
        r.font.color.rgb = color
    return sp

def para(shape, text, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         italic=False, space_before=0, space_after=4, first=False):
    tf = shape.text_frame
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align; p.space_before = Pt(space_before); p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb = color
    return p

def icon(slide, x, y, size, glyph, color=WHITE, bg=PRIMARY, radius=0.26, fsz=None):
    sp = box(slide, x, y, size, size, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
             radius=radius, anchor=MSO_ANCHOR.MIDDLE, wrap=False, ml=0, mr=0, mt=0, mb=0)
    p = sp.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = glyph
    r.font.name = ICF; r.font.size = Pt(fsz or size*42); r.font.color.rgb = color
    return sp

def card(slide, x, y, w, h, heading, body, accent=PRIMARY, hsize=12.5, bsize=10.2, ic=None):
    box(slide, x, y, w, h, fill=WHITE, line=GREY_BD, line_w=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    box(slide, x, y, w, 0.09, fill=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    if ic:
        icon(slide, x+0.16, y+0.2, 0.42, G(ic), bg=accent)
        box(slide, x+0.66, y+0.16, w-0.8, 0.52, text=heading, size=hsize, color=NAVY,
            bold=True, anchor=MSO_ANCHOR.MIDDLE)
        box(slide, x+0.18, y+0.76, w-0.36, h-0.86, text=body, size=bsize, color=MUTE)
    else:
        box(slide, x+0.16, y+0.18, w-0.32, 0.4, text=heading, size=hsize, color=NAVY, bold=True)
        box(slide, x+0.16, y+0.6, w-0.32, h-0.72, text=body, size=bsize, color=MUTE)

def stat(slide, x, y, w, number, label, color=PRIMARY, nsize=25, lsize=9.5, h=1.15):
    b = box(slide, x, y, w, h, fill=WHITE, line=GREY_BD, line_w=1.0,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09, anchor=MSO_ANCHOR.MIDDLE)
    para(b, number, size=nsize, color=color, bold=True, align=PP_ALIGN.CENTER, first=True, space_after=1)
    para(b, label, size=lsize, color=MUTE, align=PP_ALIGN.CENTER, space_after=0)
    return b

def pill(slide, x, y, w, h, text, fill=PRIMARY, tcolor=WHITE, size=10):
    return box(slide, x, y, w, h, text=text, size=size, color=tcolor, bold=True,
               align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, fill=fill,
               shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5, wrap=False)

def eyebrow(slide, text, x=CX, y=0.88, color=PRIMARY):
    return box(slide, x, y, 7.0, 0.26, text=text.upper(), size=10.5, color=color,
               bold=True, anchor=MSO_ANCHOR.MIDDLE)

def footer_note(slide, text, color=MUTE):
    return box(slide, CX, 5.18, CW, 0.28, text=text, size=8.5, color=color,
               italic=True, anchor=MSO_ANCHOR.MIDDLE)

def list_row(slide, x, y, w, glyph, head, sub, accent=PRIMARY, h=0.66, hsize=11.5, ssize=9.3):
    icon(slide, x, y+0.02, 0.4, glyph, bg=accent)
    box(slide, x+0.54, y-0.04, w-0.54, 0.3, text=head, size=hsize, color=NAVY, bold=True,
        anchor=MSO_ANCHOR.BOTTOM, mb=0)
    box(slide, x+0.54, y+0.26, w-0.54, h-0.24, text=sub, size=ssize, color=MUTE, mt=0)

def bullets(slide, x, y, w, h, items, size=11, color=INK, marker="▸", mcolor=None, gap=5):
    mcolor = mcolor or PRIMARY
    sp = box(slide, x, y, w, h, ml=0.02, mr=0.04, mt=0.02)
    tf = sp.text_frame
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        rm = p.add_run(); rm.text = marker + "  "
        rm.font.name=FONT; rm.font.size=Pt(size); rm.font.bold=True; rm.font.color.rgb=mcolor
        if isinstance(it, tuple):
            r1=p.add_run(); r1.text=it[0]; r1.font.name=FONT; r1.font.size=Pt(size); r1.font.bold=True; r1.font.color.rgb=color
            r2=p.add_run(); r2.text=it[1]; r2.font.name=FONT; r2.font.size=Pt(size); r2.font.color.rgb=color
        else:
            r=p.add_run(); r.text=it; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=color
    return sp

# ===================================================================
os.makedirs(os.path.dirname(OUT), exist_ok=True)
shutil.copy(SRC, OUT)
prs = Presentation(OUT)
L_DIV, L_STD = prs.slide_layouts[2], prs.slide_layouts[4]
xml_slides = prs.slides._sldIdLst
for sid in list(xml_slides):
    try: prs.part.drop_rel(sid.get(qn('r:id')))
    except Exception: pass
    xml_slides.remove(sid)

def content(title, kicker=None):
    s = prs.slides.add_slide(L_STD); _set_title(s, title); _blank_footer(s)
    if kicker: eyebrow(s, kicker)
    return s

def dark():
    s = prs.slides.add_slide(L_DIV); _blank_footer(s)
    return s

# --------------------------------------------------------- 1. COVER
s = dark()
box(s, 0.7, 1.16, 8.6, 0.34, text="LIVE WEBINAR  ·  20 MINUTES", size=13, color=SECOND, bold=True)
box(s, 0.66, 1.52, 8.7, 1.5, text="AccuKnox SIEM", size=42, color=WHITE, bold=True)
box(s, 0.7, 2.72, 8.4, 0.6,
    text="Raw alerts do not stop attacks. Correlated intelligence across your whole cloud does.",
    size=14, color=NAVY_TXT)
box(s, 0.7, 3.62, 0.14, 0.92, fill=PRIMARY)
box(s, 0.96, 3.58, 4.0, 0.5, text="Aditya Raj", size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM, mb=0)
box(s, 0.96, 4.02, 4.0, 0.34, text="Sr. Security Engineer  ·  Speaker", size=10.5, color=NAVY_TXT, mt=0)
box(s, 5.30, 3.58, 4.0, 0.5, text="Atharva Shah", size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM, mb=0)
box(s, 5.30, 4.02, 4.0, 0.34, text="Technical Content Lead & DevRel  ·  Moderator", size=10.5, color=NAVY_TXT, mt=0)

# --------------------------------------------------------- 2. AGENDA
s = content("Five questions in twenty minutes", "How this runs")
rows = [
 ("bulb", "What AccuKnox SIEM is", "A SaaS log platform, separate from CNAPP, natively wired into it.", PRIMARY),
 ("cloud", "What it ingests, and how much", "Cloud sinks, vendor APIs, syslog. 100K events a second.", NAVY),
 ("search", "How detection works", "Detectors, correlation, threat intel, MITRE ATT&CK.", PURPLE),
 ("cert", "Retention and audit", "One year retained, six months searchable, export on demand.", GREEN_DK),
 ("gears", "SIEM plus CNAPP plus AskADA", "One console for posture and detections, with a GenAI copilot.", RED),
]
for i, (g, hd, sub, ac) in enumerate(rows):
    list_row(s, CX+0.06, 1.36+i*0.74, CW-0.2, G(g), str(i+1)+".  "+hd, sub, accent=ac)
footer_note(s, "Questions in the chat as we go. We take what fits at the end and answer the rest by email.")

# --------------------------------------------------------- 3. Q1 WHAT IT IS
s = content("A SIEM built for teams that already run cloud security", "Question 1  ·  What it is")
c = [
 ("cloud", "Cloud native, delivered as SaaS", PRIMARY,
  "Runs at siem.accuknox.com. Nothing to rack, nothing to size. A separate product from AccuKnox CNAPP."),
 ("swap" if 'swap' in IC else "gears", "Native link into CNAPP", NAVY,
  "SIEM detections and threat analysis alerts surface inside the CNAPP console, next to posture findings."),
 ("gear", "One schema across every source", PURPLE,
  "Cloud, SaaS, endpoint and network events are normalized on the way in, then correlated as one stream."),
 ("bolt", "Detection first, by design", GREEN_DK,
  "Ingestion, detection, dashboards, alerting and retention. Remediation stays with your own response process."),
]
cw = (CW-0.3)/2; ch = 1.62
for i, (g, hd, ac, bd) in enumerate(c):
    card(s, CX+(i % 2)*(cw+0.3), 1.32+(i//2)*(ch+0.22), cw, ch, hd, bd, accent=ac, hsize=12.5, bsize=10.4, ic=g)

# --------------------------------------------------------- 4. Q1b THE FLOW
s = content("Sources in, normalized, detected, routed out", "Question 1  ·  The flow")
steps = [
 ("Sources", "AWS, GCP, OCI, SaaS APIs,\nEDR, syslog appliances", PRIMARY),
 ("Ingestion", "LogStash, Bulk API,\nData Prepper", NAVY),
 ("Detection", "Detectors, correlation,\nthreat intel, ATT&CK", PURPLE),
 ("Output", "Dashboards, alerts,\nCNAPP, Slack, SNS", GREEN_DK),
]
bw = 2.06; gap = 0.32
for i, (hd, bd, ac) in enumerate(steps):
    x = CX + i*(bw+gap)
    box(s, x, 1.60, bw, 1.60, fill=WHITE, line=GREY_BD, line_w=1.0,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    box(s, x, 1.60, bw, 0.09, fill=ac, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    box(s, x+0.12, 1.80, bw-0.24, 0.38, text=hd, size=13, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    box(s, x+0.12, 2.22, bw-0.24, 0.90, text=bd, size=10, color=MUTE, align=PP_ALIGN.CENTER)
    if i < 3:
        box(s, x+bw+0.04, 2.24, gap-0.08, 0.32, text="→", size=16, color=GREY_BD,
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
box(s, CX, 3.50, CW, 1.30, fill=RGBColor(0xEC,0xEC,0xFB), line=None,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
box(s, CX+0.24, 3.66, CW-0.48, 0.34, text="Where CNAPP fits", size=12, color=NAVY, bold=True)
box(s, CX+0.24, 3.98, CW-0.48, 0.72,
    text="CNAPP calls SIEM APIs for CloudTrail and cloud access data. KubeArmor runtime events flow the "
         "other way as a custom SIEM log type, with detectors written for them. One asset and findings view covers both.",
    size=10.6, color=MUTE)

# --------------------------------------------------------- 5. Q2 INGESTION
s = content("Three routes in, and the numbers behind them", "Question 2  ·  Ingestion and scale")
routes = [
 ("cloud", "Native cloud sinks", PRIMARY,
  "AWS CloudTrail, VPC Flow, WAF, Config, Route 53, Network Firewall. GCP Cloud Logging sinks. OCI Logging."),
 ("chip", "Vendor APIs", NAVY,
  "Google Workspace, Slack, Jamf, MongoDB Atlas, CrowdStrike FDR, AquaSec. No forwarder to babysit."),
 ("server", "Syslog middleware", PURPLE,
  "Palo Alto and Cisco Meraki forward syslog to a VM that converts to JSON, then ships via Bulk API."),
]
cw3 = (CW-0.44)/3
for i, (g, hd, ac, bd) in enumerate(routes):
    card(s, CX+i*(cw3+0.22), 1.32, cw3, 1.72, hd, bd, accent=ac, hsize=12.5, bsize=10.2, ic=g)
sw = (CW-0.66)/4
figs = [("100K", "events per second", PRIMARY), ("1 TB", "per day, per tenant", NAVY),
        ("15+", "source types onboarded", PURPLE), ("1", "isolated pipeline per tenant", GREEN_DK)]
for i, (n, l, c_) in enumerate(figs):
    stat(s, CX+i*(sw+0.22), 3.30, sw, n, l, color=c_, nsize=24, lsize=9.2, h=1.10)
footer_note(s, "Connectors cover HTTP, Kafka, OTel, S3, Kinesis, DynamoDB, DocumentDB and Fluentd inputs.")

# --------------------------------------------------------- 6. Q3 DETECTION
s = content("What turns a log line into an alert worth reading", "Question 3  ·  Detection")
left = [
 ("Out of the box detectors. ", "AWS CloudTrail, GCP and Kubernetes ship ready to run."),
 ("Custom detectors. ", "Written per customer across syslog, KubeArmor, EDR, CloudTrail and O365 log types."),
 ("Multivector correlation. ", "Cloud, identity and network events analyzed together, not source by source."),
 ("Anomaly detection and hunting. ", "Analysts query the indexed window directly when they have a hypothesis."),
]
box(s, CX, 1.32, 4.44, 0.34, text="How detections are produced", size=12, color=NAVY, bold=True)
bullets(s, CX, 1.68, 4.44, 2.20, left, size=10.6, gap=7)
box(s, 5.16, 1.32, 4.44, 0.34, text="What gets added on top", size=12, color=NAVY, bold=True)
right = [
 ("Threat intel enrichment. ", "IP and domain reputation from external sources, including AlienVault."),
 ("MITRE ATT&CK mapping. ", "Cloud, container and endpoint tactics, not cloud API calls alone."),
 ("Prioritization. ", "Correlation exists to cut analyst hours spent on noise, not to raise alert volume."),
 ("Routing. ", "Slack, webhooks and Amazon SNS, plus Jira, ServiceNow or email for ticketing."),
]
bullets(s, 5.16, 1.68, 4.44, 2.20, right, size=10.6, gap=7, mcolor=PURPLE)
box(s, CX, 4.10, CW, 0.76, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
box(s, CX+0.28, 4.10, CW-0.56, 0.76,
    text="No automated remediation is configured. Alerts are informational and route into your own response process.",
    size=11.5, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

# --------------------------------------------------------- 7. Q4 RETENTION
s = content("One year retained, six months searchable", "Question 4  ·  Retention and audit")
box(s, CX, 1.34, 5.86, 0.34, text="How the storage model works", size=12, color=NAVY, bold=True)
box(s, CX, 1.74, 2.84, 1.06, fill=WHITE, line=GREY_BD, line_w=1.0,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
box(s, CX+0.14, 1.86, 2.56, 0.34, text="Months 1 to 6", size=11.5, color=PRIMARY, bold=True)
box(s, CX+0.14, 2.18, 2.56, 0.56, text="Metadata indexed and fully searchable", size=10, color=MUTE)
box(s, CX+3.02, 1.74, 2.84, 1.06, fill=WHITE, line=GREY_BD, line_w=1.0,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
box(s, CX+3.16, 1.86, 2.56, 0.34, text="Months 7 to 12", size=11.5, color=PURPLE, bold=True)
box(s, CX+3.16, 2.18, 2.56, 0.56, text="Cold S3 compatible archive, restorable on demand", size=10, color=MUTE)
ret = [
 ("Indexing at ingestion. ", "Source, timestamp, event type, severity, account and region, which is what keeps six month search fast."),
 ("Audit export. ", "Scheduled or on demand, in JSON or CSV."),
 ("Quotas. ", "Per tenant ingestion limits and storage caps keep the bill predictable."),
 ("Snapshots. ", "Periodic snapshots of indexed data for recovery and continuity."),
]
bullets(s, CX, 2.96, 5.86, 1.90, ret, size=10.6, gap=7)
box(s, 6.50, 1.34, 3.10, 0.34, text="Frameworks it lines up with", size=12, color=NAVY, bold=True)
for i, fw in enumerate(["ISO/IEC 27001", "SOC 2", "PCI DSS", "GDPR", "NIST SP 800-207"]):
    pill(s, 6.50, 1.76+i*0.56, 3.10, 0.44, fw, fill=RGBColor(0xEC,0xEC,0xFB), tcolor=NAVY, size=11)

# --------------------------------------------------------- 8. Q5 CNAPP
s = content("One console for posture and detections", "Question 5  ·  SIEM plus CNAPP")
pairs = [
 ("cloud", "CNAPP pulls from SIEM", PRIMARY,
  "CNAPP calls SIEM APIs for CloudTrail and cloud access data, so cloud activity sits beside posture findings."),
 ("shield", "KubeArmor pushes into SIEM", NAVY,
  "Runtime events land as a custom SIEM log type with purpose built detectors written against them."),
 ("eye", "Alerts triaged in one place", PURPLE,
  "Filter by module, status, log source and severity. Every alert opens to the raw payload for audit."),
 ("chat", "Routed where the team works", GREEN_DK,
  "Slack, webhooks and Amazon SNS carry detections into the SOC workflow you already run."),
]
for i, (g, hd, ac, bd) in enumerate(pairs):
    card(s, CX+(i % 2)*(cw+0.3), 1.32+(i//2)*(ch+0.22), cw, ch, hd, bd, accent=ac, hsize=12.5, bsize=10.4, ic=g)

# --------------------------------------------------------- 9. Q5b ASKADA
s = content("AskADA reads both datasets at once", "Question 5  ·  The GenAI copilot")
box(s, CX, 1.36, CW, 1.06, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
box(s, CX+0.30, 1.36, CW-0.60, 1.06,
    text="AskADA is a GenAI copilot running on top of SIEM and CNAPP data together. It combines your own findings "
         "and telemetry with external threat intelligence, so analysts investigate across cloud and endpoint assets "
         "without pivoting between tools.",
    size=12.5, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
outs = [
 ("bolt", "Faster triage", "Multivector threat analysis runs against your data, not a generic model prompt.", PRIMARY),
 ("list", "Asset inventory in the loop", "CNAPP's inventory and findings database feed the same analysis.", PURPLE),
 ("person", "Fewer tools per analyst", "One question spans cloud logs, runtime events and posture findings.", GREEN_DK),
]
for i, (g, hd, bd, ac) in enumerate(outs):
    card(s, CX+i*(cw3+0.22), 2.66, cw3, 1.72, hd, bd, accent=ac, hsize=12.5, bsize=10.2, ic=g)

# --------------------------------------------------------- 10. RECAP
s = content("Five things to take away", "Recap")
recap = [
 ("check", "SaaS SIEM, separate from CNAPP, natively linked to it", PRIMARY),
 ("check", "Three ingestion routes: cloud sinks, vendor APIs, syslog middleware", NAVY),
 ("check", "100K events per second and 1TB a day per tenant, isolated per tenant", PURPLE),
 ("check", "One year retention, six months searchable, JSON or CSV audit export", GREEN_DK),
 ("check", "AskADA puts a GenAI copilot over SIEM and CNAPP data together", RED),
]
for i, (g, txt, ac) in enumerate(recap):
    icon(s, CX+0.06, 1.42+i*0.72, 0.42, G(g), bg=ac)
    box(s, CX+0.62, 1.40+i*0.72, CW-0.7, 0.46, text=txt, size=13, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# --------------------------------------------------------- 11. NEXT STEPS
s = content("Three ways to take this further", "Next steps")
ns = [
 ("rocket", "Book a POC", PRIMARY,
  "A phased rollout: log source onboarding, then detection tuning, then dashboards and compliance reporting. "
  "You get a scoped timeline and a checklist before anything connects."),
 ("books", "Read the onboarding guides", NAVY,
  "AWS and GCP SIEM onboarding are documented step by step on AccuKnox help docs. AWS runs off a CloudFormation template."),
 ("doc", "Ask for the full proposal", PURPLE,
  "Architecture, the source by source ingestion map, the phased roadmap, and the retention model in detail."),
]
for i, (g, hd, ac, bd) in enumerate(ns):
    card(s, CX+i*(cw3+0.22), 1.36, cw3, 2.10, hd, bd, accent=ac, hsize=13, bsize=10.4, ic=g)
box(s, CX, 3.76, CW, 1.08, fill=RGBColor(0xEC,0xEC,0xFB), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
box(s, CX+0.28, 3.88, CW-0.56, 0.34, text="Going out to everyone by email", size=12, color=NAVY, bold=True)
box(s, CX+0.28, 4.20, CW-0.56, 0.56,
    text="Slides, the recording, both onboarding guides, and a link to book time with the SIEM team.",
    size=11, color=MUTE)

# --------------------------------------------------------- 12. CLOSE
s = dark()
box(s, 0.7, 1.50, 8.6, 0.34, text="THANK YOU", size=13, color=SECOND, bold=True)
box(s, 0.66, 1.86, 8.7, 1.10, text="See AccuKnox SIEM in action", size=34, color=WHITE, bold=True)
box(s, 0.7, 2.94, 8.4, 0.44, text="Book a walkthrough with the team that built it.", size=14, color=NAVY_TXT)
pill(s, 0.7, 3.56, 3.05, 0.56, "support@accuknox.com", fill=PRIMARY, size=12)
pill(s, 3.95, 3.56, 3.05, 0.56, "accuknox.com/demo", fill=PURPLE, size=12)
box(s, 0.7, 4.42, 8.4, 0.40,
    text="Aditya Raj, Sr. Security Engineer   ·   Atharva Shah, Technical Content Lead & DevRel",
    size=10.5, color=NAVY_TXT)

prs.save(OUT)
print("saved", os.path.abspath(OUT), len(prs.slides.__iter__.__self__._sldIdLst), "slides")
