# Polish the AccuKnox branding on D:\AccuKnox_OnPrem_Security_Assessment_Revised.docx.
# Content is left untouched. Only three brand-fidelity fixes are applied:
#   1. Cover title -> bold Primary Blue (#0046FF), Space Grotesk (was thin/black).
#   2. Header logo -> swap in the official high-res assets/logos/accuknox-logo-light-bg.png.
#   3. Header label -> collapse the stray double space in "Security Assessment  Report".
import os, shutil, zipfile
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = r"D:\AccuKnox_OnPrem_Security_Assessment_Revised.docx"
OUT = os.path.join(REPO, "output", "AccuKnox_OnPrem_Security_Assessment_Branded.docx")
OFFICIAL_LOGO = os.path.join(REPO, "assets", "logos", "accuknox-logo-light-bg.png")

PRIMARY = "0046FF"
BODY_FONT = "Space Grotesk"

os.makedirs(os.path.dirname(OUT), exist_ok=True)

doc = Document(SRC)


def set_run_font(run, name=BODY_FONT):
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)


# ---- 1. Cover title: bold Primary Blue, brand font, keep the 36pt from the style
title_p = doc.paragraphs[0]
assert title_p.style.name == "Title", f"expected Title, got {title_p.style.name}"
assert "AccuKnox Security Assessment" in title_p.text, title_p.text
for run in title_p.runs:
    if run.text.strip():
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(PRIMARY)
        set_run_font(run)

# ---- 3. Header label: collapse the double space (no wording change)
for section in doc.sections:
    for p in section.header.paragraphs:
        for run in p.runs:
            if "Security Assessment" in run.text and "  " in run.text:
                run.text = " ".join(run.text.split())

doc.save(OUT)

# ---- 2. Swap the header logo bytes for the official high-res asset.
# The single media file (word/media/image1.png) is the header logo; the drawing
# carries a fixed extent, so the higher-res image displays at the same size.
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    logo_bytes = open(OFFICIAL_LOGO, "rb").read()
    for item in zin.infolist():
        data = logo_bytes if item.filename == "word/media/image1.png" else zin.read(item.filename)
        zout.writestr(item, data)
shutil.move(tmp, OUT)

print("wrote", OUT)
